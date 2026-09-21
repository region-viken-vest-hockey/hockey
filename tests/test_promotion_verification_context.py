"""Promotion must be bound to the reviewed Stage 4 verification context.

The regression these tests lock in: promotion used to rebuild a planning
problem from whatever Stage 1/2 state was present (or fall back to a
context-free verifier), so an input-constrained schedule that Stage 4 had
accepted could be rejected with false hard violations -- and, worse, a later
run's config/calendar evidence could silently change what "verified" meant.
"""

from __future__ import annotations

import itertools
import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from tournament_scheduler.pipeline.fingerprints import stable_payload_sha256
from tournament_scheduler.pipeline.run_manifest import RunManifest
from tournament_scheduler.pipeline.stage4_export import run as run_export
from tournament_scheduler.pipeline.state import PipelineState, StageName, StageStatus
from tournament_scheduler.pipeline.verification_context import (
    VerificationContextError,
    resolve_promotion_verification_context,
)
from tournament_scheduler.planning_contract import verify_candidate
from tournament_scheduler.season_state import (
    SeasonStateError,
    decisions_path,
    promote_from_stage3,
    schedule_path,
)
from tournament_scheduler.testing.reviewed_export import write_reviewed_stage4_export


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _write_input_workbook(path: Path, raw: dict) -> None:
    workbook = openpyxl.Workbook()
    settings = workbook.active
    settings.title = "Innstillinger"
    settings.append(["felt", "verdi"])
    for key in ("start_date", "end_date"):
        settings.append([key, raw[key]])
    age_groups = workbook.create_sheet("Aldersgrupper")
    age_groups.append(["age_group", "parallel_games", "round_length_minutes"])
    for age_group in raw.get("age_groups", []):
        age_groups.append([age_group, raw.get("parallel_games", {}).get(age_group), None])
    teams = workbook.create_sheet("Lag")
    teams.append(["club", "label", "age_group"])
    for team in raw.get("teams", []):
        teams.append([team["club"], team["label"], team["age_group"]])
    sources = workbook.create_sheet("Kilder")
    sources.append(["name", "type", "url"])
    workbook.save(path)


def _round_robin_games(labels: list[str]) -> list[dict]:
    return [
        {"home": home, "away": away, "parallel_slot": 0, "round_number": 1 + index // 2}
        for index, (home, away) in enumerate(itertools.combinations(labels, 2))
    ]


def _odd_team_plan() -> dict:
    """Production-shaped JU8 tournament with an input-constrained odd 5-team pool."""
    teams = [{"club": club, "label": f"JU8-{club}", "age_group": "JU8"} for club in "ABCDE"]
    return {
        "start_date": "2026-09-01",
        "end_date": "2027-04-30",
        "tournaments": [
            {
                "id": "ju8-a-20260912",
                "date": "2026-09-12",
                "arena": "Arena A",
                "age_group": "JU8",
                "host_club": "A",
                "teams": teams,
                "games": _round_robin_games([team["label"] for team in teams]),
                "start_time": "10:00",
            }
        ],
    }


def _stage_odd_team_export(tmp_path: Path, *, start_manifest: bool = False):
    """Run the real Stage 4 export over an input-constrained odd-team plan."""
    work_dir = tmp_path / ".pipeline"
    root = tmp_path / "season"
    state = PipelineState(work_dir)
    if start_manifest:
        RunManifest(work_dir).start_run("review run", input_fingerprint={})
    teams = [{"club": club, "label": f"JU8-{club}", "age_group": "JU8"} for club in "ABCDE"]
    input_path = tmp_path / "input.xlsx"
    _write_input_workbook(
        input_path,
        {"start_date": "2026-09-01", "end_date": "2027-04-30", "age_groups": ["JU8"], "teams": teams},
    )
    state.write_stage(
        StageName.CONFIG,
        {
            "input_path": str(input_path),
            "teams": teams,
            "round_length_minutes": {"JU8": 10},
            "ice_time_minutes": {"JU8": 120},
        },
        status=StageStatus.DONE,
    )
    plan = _odd_team_plan()
    state.write_stage(StageName.PLANNING, {"plan": plan}, status=StageStatus.DONE)
    result = run_export(
        {"plan": plan}, state, export_dir=str(tmp_path / "export"), timestamped_export=False
    )
    return work_dir, root, state, plan, result


# ---------------------------------------------------------------------------
# Immediate production defect
# ---------------------------------------------------------------------------


def test_stage4_accepts_input_constrained_odd_team_plan_and_promotion_succeeds(tmp_path):
    work_dir, root, _state, plan, result = _stage_odd_team_export(tmp_path)

    # Stage 4 accepted the exact reviewed candidate with its real planning problem.
    assert result["errors"] == []
    assert result["verify_result"]["ok"], result["verify_result"]["violations"]

    # The context-free verifier would have produced the false hard violation
    # that made promotion fail in production.
    context_free = verify_candidate(plan)
    assert context_free["ok"] is False
    assert "bye_team_not_allowed" in {v["code"] for v in context_free["violations"]}

    schedule, _decisions = promote_from_stage3(work_dir=work_dir, root=root, actor="tester")

    assert schedule["season"] == "2026-2027"
    assert schedule["plan"]["tournaments"][0]["id"] == "ju8-a-20260912"


def test_same_run_resume_remains_promotable(tmp_path):
    work_dir, root, _state, _plan, result = _stage_odd_team_export(tmp_path, start_manifest=True)

    run_id = RunManifest(work_dir).read()["run_id"]
    assert result["verification_context"]["run_id"] == run_id

    schedule, _decisions = promote_from_stage3(work_dir=work_dir, root=root, actor="tester")
    assert schedule["promoted_from"]["run_id"] == run_id


def test_promotion_persists_reviewed_export_provenance(tmp_path):
    work_dir, root, _state, _plan, result = _stage_odd_team_export(tmp_path)

    schedule, _decisions = promote_from_stage3(work_dir=work_dir, root=root, actor="tester")
    promoted_from = schedule["promoted_from"]

    assert promoted_from["run_id"] == "legacy"
    assert promoted_from["stage4_export_fingerprint"] == result["export_fingerprint"]
    assert promoted_from["verification_context_problem_fingerprint"] == result["verification_context"][
        "problem_fingerprint"
    ]
    assert promoted_from["verification_context_candidate_fingerprint"] == result["export_fingerprint"]
    assert promoted_from["verification_context_verified_ok"] is True
    assert schedule["verification_context"]["problem_fingerprint"] == result["verification_context"]["problem_fingerprint"]


def test_promotion_uses_reconciled_stage4_plan_snapshot_not_stale_stage3_fields(tmp_path):
    work_dir, root, state, plan, _result = _stage_odd_team_export(tmp_path)

    stage3_plan = dict(plan)
    stage3_plan["publication_readiness"] = {"status": "stale_stage3"}
    state.write_stage(StageName.PLANNING, {"plan": stage3_plan}, status=StageStatus.DONE)

    reviewed_plan = dict(plan)
    reviewed_plan["publication_readiness"] = {"status": "stage4_reconciled"}
    run_export({"plan": reviewed_plan}, state, export_dir=str(tmp_path / "export2"), timestamped_export=False)

    schedule, _decisions = promote_from_stage3(work_dir=work_dir, root=root, actor="tester")

    assert schedule["plan"]["publication_readiness"] == {"status": "stage4_reconciled"}


# ---------------------------------------------------------------------------
# Stale-state protection
# ---------------------------------------------------------------------------


def test_new_run_id_blocks_promotion(tmp_path):
    work_dir, root, _state, _plan, _result = _stage_odd_team_export(tmp_path)
    # A new run starts in the same workspace after the reviewed export.
    RunManifest(work_dir).start_run("a later run", input_fingerprint={})

    with pytest.raises(SeasonStateError) as excinfo:
        promote_from_stage3(work_dir=work_dir, root=root, actor="tester")
    assert "verification context belongs to run" in str(excinfo.value)

    assert not schedule_path("2026-2027", root=root).exists()
    assert not decisions_path("2026-2027", root=root).exists()


def test_stage3_candidate_fingerprint_mismatch_blocks_promotion(tmp_path):
    work_dir, root, state, plan, _result = _stage_odd_team_export(tmp_path)
    # Overwrite the Stage 3 checkpoint out-of-band (a legacy/out-of-band path
    # that does not trigger downstream stale-marking) with a different
    # candidate after review.
    planning_path = state.checkpoint_path(StageName.PLANNING)
    envelope = json.loads(planning_path.read_text(encoding="utf-8"))
    envelope["data"]["plan"]["tournaments"] = [
        {**plan["tournaments"][0], "date": "2026-09-19"},
    ]
    planning_path.write_text(json.dumps(envelope), encoding="utf-8")

    with pytest.raises(SeasonStateError) as excinfo:
        promote_from_stage3(work_dir=work_dir, root=root, actor="tester")
    assert "no longer matches the reviewed Stage 4 export" in str(excinfo.value)


def test_missing_verification_context_blocks_instead_of_context_free_fallback(tmp_path):
    work_dir = tmp_path / ".pipeline"
    root = tmp_path / "season"
    state = PipelineState(work_dir)
    plan = _odd_team_plan()
    state.write_stage(StageName.PLANNING, {"plan": plan}, status=StageStatus.DONE)
    # A legacy export with no verification-context provenance.
    state.write_stage(
        StageName.EXPORT,
        {
            "export_dir": str(tmp_path / "export"),
            "output_files": {},
            "errors": [],
            "verify_result": verify_candidate(plan),
            "export_fingerprint": stable_payload_sha256(plan["tournaments"]),
        },
        status=StageStatus.DONE,
    )

    with pytest.raises(SeasonStateError) as excinfo:
        promote_from_stage3(work_dir=work_dir, root=root, actor="tester")
    assert "no verification-context provenance" in str(excinfo.value)
    assert not schedule_path("2026-2027", root=root).exists()


def test_missing_reviewed_plan_snapshot_blocks_promotion(tmp_path):
    work_dir, root, state, _plan, _result = _stage_odd_team_export(tmp_path)
    export_path = state.checkpoint_path(StageName.EXPORT)
    envelope = json.loads(export_path.read_text(encoding="utf-8"))
    envelope["data"].pop("reviewed_plan")
    export_path.write_text(json.dumps(envelope), encoding="utf-8")

    with pytest.raises(SeasonStateError) as excinfo:
        promote_from_stage3(work_dir=work_dir, root=root, actor="tester")
    assert "no reviewed final plan snapshot" in str(excinfo.value)


def test_missing_problem_blocks_promotion(tmp_path):
    work_dir = tmp_path / ".pipeline"
    root = tmp_path / "season"
    state = PipelineState(work_dir)
    plan = _odd_team_plan()
    state.write_stage(StageName.PLANNING, {"plan": plan}, status=StageStatus.DONE)
    fingerprint = stable_payload_sha256(plan["tournaments"])
    state.write_stage(
        StageName.EXPORT,
        {
            "verify_result": {"ok": True, "violations": []},
            "export_fingerprint": fingerprint,
            "verification_context": {
                "schema_version": 1,
                "run_id": "legacy",
                "candidate_fingerprint": fingerprint,
                "problem": None,
                "problem_fingerprint": None,
                "verify_ok": True,
            },
        },
        status=StageStatus.DONE,
    )

    with pytest.raises(SeasonStateError) as excinfo:
        promote_from_stage3(work_dir=work_dir, root=root, actor="tester")
    assert "no normalized verification problem" in str(excinfo.value)


def test_promotion_uses_bound_context_even_if_stage1_config_changes(tmp_path):
    work_dir, root, state, _plan, _result = _stage_odd_team_export(tmp_path)

    # Replace the Stage 1 config out-of-band (without starting a new run) with
    # a pool that would make the reviewed 5-team shape avoidable and therefore
    # a hard violation. Promotion must still verify against the *bound*
    # reviewed context, not silently adopt this newer config.
    config_path = state.checkpoint_path(StageName.CONFIG)
    envelope = json.loads(config_path.read_text(encoding="utf-8"))
    envelope["data"]["teams"] = [
        {"club": club, "label": f"JU8-{club}", "age_group": "JU8"} for club in "ABCDEF"
    ]
    config_path.write_text(json.dumps(envelope), encoding="utf-8")

    schedule, _decisions = promote_from_stage3(work_dir=work_dir, root=root, actor="tester")
    assert schedule["season"] == "2026-2027"


def test_returned_manifest_legacy_run_id_matches_export(tmp_path):
    """A workspace with no manifest still round-trips its synthesized run id."""
    work_dir, root, state, _plan, result = _stage_odd_team_export(tmp_path)
    assert result["verification_context"]["run_id"] == RunManifest(work_dir).read()["run_id"]

    bound = resolve_promotion_verification_context(
        work_dir=str(work_dir), candidate=state.read_stage(StageName.PLANNING)["plan"]
    )
    assert bound["problem"] is not None
    assert bound["problem_fingerprint"] == result["verification_context"]["problem_fingerprint"]


def test_tampered_problem_fingerprint_blocks_promotion(tmp_path):
    work_dir = tmp_path / ".pipeline"
    state = PipelineState(work_dir)
    plan = _odd_team_plan()
    state.write_stage(StageName.PLANNING, {"plan": plan}, status=StageStatus.DONE)
    checkpoint = write_reviewed_stage4_export(state)

    export_path = state.checkpoint_path(StageName.EXPORT)
    envelope = json.loads(export_path.read_text(encoding="utf-8"))
    envelope["data"]["verification_context"]["problem_fingerprint"] = "0" * 64
    export_path.write_text(json.dumps(envelope), encoding="utf-8")

    with pytest.raises(SeasonStateError) as excinfo:
        promote_from_stage3(work_dir=work_dir, root=tmp_path / "season", actor="tester")
    assert "does not match its recorded fingerprint" in str(excinfo.value)


def test_resolver_refuses_stale_export(tmp_path):
    work_dir = tmp_path / ".pipeline"
    state = PipelineState(work_dir)
    plan = _odd_team_plan()
    state.write_stage(StageName.PLANNING, {"plan": plan}, status=StageStatus.DONE)
    write_reviewed_stage4_export(state)

    # A later Stage 1 write marks the reviewed export stale.
    state.write_stage(StageName.CONFIG, {"teams": []}, status=StageStatus.DONE)

    with pytest.raises(VerificationContextError) as excinfo:
        resolve_promotion_verification_context(work_dir=str(work_dir), candidate=plan)
    assert "not a completed, non-stale reviewed handoff" in str(excinfo.value)


# ---------------------------------------------------------------------------
# CLI surface
# ---------------------------------------------------------------------------


def test_season_promote_cli_succeeds_for_reviewed_odd_team_export(tmp_path, capsys):
    from tournament_scheduler.cli.rvv_cli import main

    work_dir, root, _state, _plan, _result = _stage_odd_team_export(tmp_path)

    rc = main(["season", "promote", "--work-dir", str(work_dir), "--root", str(root), "--json"])

    assert rc == 0
    captured = capsys.readouterr().out
    summary = json.loads(captured[captured.index("{") : captured.rindex("}") + 1])
    assert summary["season"] == "2026-2027"
    assert summary["tournament_count"] == 1


def test_canonical_season_export_uses_durable_context_after_pipeline_deleted(tmp_path, capsys):
    from tournament_scheduler.cli.rvv_cli import main

    work_dir, root, _state, _plan, _result = _stage_odd_team_export(tmp_path)
    assert main(["season", "promote", "--work-dir", str(work_dir), "--root", str(root)]) == 0
    capsys.readouterr()
    shutil.rmtree(work_dir)

    rc = main([
        "season",
        "export",
        "--season",
        "2026-2027",
        "--work-dir",
        str(work_dir),
        "--root",
        str(root),
        "--export-dir",
        str(tmp_path / "canonical-export"),
        "--flat",
        "--json",
    ])

    assert rc == 0
    captured = capsys.readouterr().out
    result = json.loads(captured[captured.index("{") : captured.rindex("}") + 1])
    assert result["verify_result"]["ok"] is True
    assert result["canonical_season"] == "2026-2027"


def test_canonical_season_export_does_not_consume_later_stage1_config(tmp_path, capsys):
    from tournament_scheduler.cli.rvv_cli import main

    work_dir, root, state, _plan, _result = _stage_odd_team_export(tmp_path)
    assert main(["season", "promote", "--work-dir", str(work_dir), "--root", str(root)]) == 0
    capsys.readouterr()
    state.write_stage(
        StageName.CONFIG,
        {
            "start_date": "2026-09-01",
            "end_date": "2027-04-30",
            "teams": [{"club": club, "label": f"JU8-{club}", "age_group": "JU8"} for club in "ABCDEF"],
            "round_length_minutes": {"JU8": 10},
            "ice_time_minutes": {"JU8": 120},
        },
        status=StageStatus.DONE,
    )

    rc = main([
        "season",
        "export",
        "--season",
        "2026-2027",
        "--work-dir",
        str(work_dir),
        "--root",
        str(root),
        "--export-dir",
        str(tmp_path / "canonical-export"),
        "--flat",
        "--json",
    ])

    assert rc == 0
    captured = capsys.readouterr().out
    result = json.loads(captured[captured.index("{") : captured.rindex("}") + 1])
    assert result["verify_result"]["ok"] is True


def test_season_promote_cli_refuses_stale_handoff(tmp_path, capsys):
    from tournament_scheduler.cli.rvv_cli import main

    work_dir, root, _state, _plan, _result = _stage_odd_team_export(tmp_path)
    RunManifest(work_dir).start_run("a later run", input_fingerprint={})

    rc = main(["season", "promote", "--work-dir", str(work_dir), "--root", str(root)])

    assert rc == 1
    assert "verification context belongs to run" in capsys.readouterr().out
    assert not schedule_path("2026-2027", root=root).exists()
    assert not decisions_path("2026-2027", root=root).exists()
