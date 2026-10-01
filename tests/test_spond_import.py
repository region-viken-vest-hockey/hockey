"""Tests for the Spond bulk-import projection and workbooks."""

from datetime import date, time

import openpyxl

from tournament_scheduler.models import Game, SeasonPlan, Team, Tournament
from tournament_scheduler.spond.spond_exporter import SpondExporter
from tournament_scheduler.spond.spond_import import (
    HELPER_HEADERS,
    SPOND_IMPORT_HEADERS,
    SPOND_IMPORT_SHEET,
    project_spond_import_rows,
    write_per_team_workbooks,
)


def _plan(*, cancel_second: bool = False) -> SeasonPlan:
    jar1, jar2 = Team("Jar", "Jar 1", "U10"), Team("Jar", "Jar 2", "U10")
    holmen, frisk = Team("Holmen", "Holmen 1", "U10"), Team("Frisk Asker", "Frisk Asker 3", "U10")
    first = Tournament(
        date=date(2026, 11, 14),
        arena="Holmen Ishall",
        age_group="U10",
        id="rvv-0123",
        teams=[holmen, frisk, jar2],
        games=[Game(home=holmen, away=frisk, round_number=1)],
        host_club="Holmen",
        start_time="09:00",
    )
    second = Tournament(
        date=date(2026, 12, 5),
        arena="Jar Isforum",
        age_group="U10",
        id="rvv-0200",
        teams=[jar1, jar2],
        games=[Game(home=jar1, away=jar2, round_number=1)],
        host_club="Jar",
        start_time="10:00",
        cancelled=cancel_second,
        cancellation_reason="Ishall stengt" if cancel_second else None,
    )
    return SeasonPlan(tournaments=[second, first])


def test_headers_match_spond_grid_exactly():
    assert SPOND_IMPORT_HEADERS == (
        "Startdato*", "Starttidspunkt", "Oppmøte", "Sluttdato",
        "Sluttidspunkt", "Tittel*", "Beskrivelse", "Sted",
    )


def test_workbook_has_table_over_all_data_rows_with_spond_columns_first(tmp_path):
    out = tmp_path / "spond.xlsx"
    SpondExporter().export(_plan(), str(out), ice_time_for_age_group={"U10": 60})
    sheet = openpyxl.load_workbook(out)[SPOND_IMPORT_SHEET]

    header = [c.value for c in sheet[1]]
    assert header == list(SPOND_IMPORT_HEADERS + HELPER_HEADERS)
    assert len(sheet.tables) == 1
    table = next(iter(sheet.tables.values()))
    assert table.ref == f"A1:M{sheet.max_row}"
    assert sheet.freeze_panes == "A2"


def test_rows_are_per_team_and_deterministic():
    rows = project_spond_import_rows(_plan(), {"U10": 60})
    assert rows == project_spond_import_rows(_plan(), {"U10": 60})
    assert [(r.start_date, r.team) for r in rows] == [
        (date(2026, 11, 14), "Frisk Asker 3"),
        (date(2026, 11, 14), "Holmen 1"),
        (date(2026, 11, 14), "Jar 2"),
        (date(2026, 12, 5), "Jar 1"),
        (date(2026, 12, 5), "Jar 2"),
    ]


def test_selected_team_gets_one_event_per_tournament():
    rows = project_spond_import_rows(_plan(), {"U10": 60}, team="Jar 2")
    assert [r.tournament_id for r in rows] == ["rvv-0123", "rvv-0200"]
    assert len({r.tournament_id for r in rows}) == len(rows)


def test_cancelled_tournaments_are_not_importable():
    rows = project_spond_import_rows(_plan(cancel_second=True), {"U10": 60})
    assert {r.tournament_id for r in rows} == {"rvv-0123"}


def test_date_time_and_end_time_mapping_and_description():
    row = project_spond_import_rows(_plan(), {"U10": 60}, team="Holmen 1")[0]
    assert row.start_date == date(2026, 11, 14)
    assert row.start_time == time(9, 0)
    assert row.end_date == date(2026, 11, 14)
    assert row.end_time is not None and row.end_time > time(9, 0)
    assert row.meet_time is None
    assert row.venue == "Holmen Ishall"
    assert row.description == (
        "U10 • Vert: Holmen • Lag: Holmen 1, Frisk Asker 3, Jar 2 • RVV-ID: rvv-0123"
    )


def test_per_team_workbooks_contain_only_that_team(tmp_path):
    written = write_per_team_workbooks(_plan(cancel_second=True), tmp_path / "spond", {"U10": 60})

    assert sorted(written) == ["Frisk Asker 3", "Holmen 1", "Jar 2"]  # Jar 1 only had the cancelled one
    assert (tmp_path / "spond" / "Jar 2.xlsx").exists()
    sheet = openpyxl.load_workbook(tmp_path / "spond" / "Jar 2.xlsx")[SPOND_IMPORT_SHEET]
    data = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(data) == 1
    assert data[0][12] == "rvv-0123"
    assert data[0][8] == "Jar 2"
