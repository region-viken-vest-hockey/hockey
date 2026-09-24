"""Canonical ``season export`` must retain the promoted public/source context.

Regression: canonical export used to call Stage 4 with
``use_pipeline_metadata=False`` and therefore rendered ``0 kilder · 0 hendelser``
and dropped the scraped-calendar/registered-team companion pages, because
canonical safety (don't trust stale ``.pipeline`` checkpoints for verification)
and public presentation metadata were coupled through one switch.

These tests promote a reviewed Stage 4 handoff that *has* source/public
metadata, then delete the entire ``.pipeline`` workspace and prove the canonical
export still rebuilds the same public bundle purely from promoted state.
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from tournament_scheduler.pipeline.cache_manager import ScrapedDataCache
from tournament_scheduler.pipeline.public_export_context import (
    PublicExportContextError,
    fingerprint_public_export_context,
    resolve_promoted_public_export_context,
    verify_public_export_context,
)
from tournament_scheduler.pipeline.run_manifest import RunManifest
from tournament_scheduler.pipeline.stage4_export import run as run_export
from tournament_scheduler.pipeline.state import PipelineState, StageName, StageStatus
from tournament_scheduler.season_state import load_export_context, promote_from_stage3


CLUBS = ("Alfa", "Bravo", "Charlie", "Delta")


def _write_input_workbook(path: Path, *, with_activities: bool = True) -> None:
    workbook = openpyxl.Workbook()
    settings = workbook.active
    settings.title = "Innstillinger"
    settings.append(["felt", "verdi"])
    settings.append(["start_date", "2026-09-01"])
    settings.append(["end_date", "2027-04-30"])
    age_groups = workbook.create_sheet("Aldersgrupper")
    age_groups.append(["age_group", "parallel_games", "round_length_minutes"])
    age_groups.append(["JU8", 2, 10])
    teams = workbook.create_sheet("Lag")
    teams.append(["club", "label", "age_group"])
    for club in CLUBS:
        teams.append([club, f"JU8-{club}", "JU8"])
    sources = workbook.create_sheet("Kilder")
    sources.append(["name", "type", "url"])
    sources.append(["Alfa", "ical", "https://example.invalid/alfa.ics"])
    if with_activities:
        activities = workbook.create_sheet("Aktiviteter")
        activities.append(["Måned", "Dato", "Aktivitet", "Sted"])
        activities.append(["Januar", 17, "Spillerutvikling JU8", "Sandefjord"])
    workbook.save(path)


def _round_robin_games(labels: list[str]) -> list[dict]:
    """Return a no-bye round-robin schedule for an even team count."""
    rotation = list(labels)
    half = len(rotation) // 2
    games: list[dict] = []
    for round_index in range(len(rotation) - 1):
        for slot in range(half):
            home = rotation[slot]
            away = rotation[len(rotation) - 1 - slot]
            # Alternate home/away by round to avoid a fixed-side streak.
            if round_index % 2:
                home, away = away, home
            games.append(
                {
                    "home": home,
                    "away": away,
                    "parallel_slot": slot,
                    "round_number": round_index + 1,
                }
            )
        rotation = [rotation[0], rotation[-1], *rotation[1:-1]]
    return games


def _plan() -> dict:
    labels = [f"JU8-{club}" for club in CLUBS]
    return {
        "start_date": "2026-09-01",
        "end_date": "2027-04-30",
        "tournaments": [
            {
                "id": "ju8-alfa-20260912",
                "date": "2026-09-12",
                "arena": "Alfa Arena",
                "age_group": "JU8",
                "host_club": "Alfa",
                "teams": [
                    {"club": club, "label": f"JU8-{club}", "age_group": "JU8"} for club in CLUBS
                ],
                "games": _round_robin_games(labels),
                "start_time": "10:00",
            }
        ],
    }


def _seed_workspace(tmp_path: Path) -> PipelineState:
    work_dir = tmp_path / ".pipeline"
    state = PipelineState(work_dir)
    RunManifest(work_dir).start_run("canonical export context run", input_fingerprint={})
    input_path = tmp_path / "input.xlsx"
    _write_input_workbook(input_path)
    state.write_stage(
        StageName.CONFIG,
        {
            "input_path": str(input_path),
            "start_date": "2026-09-01",
            "end_date": "2027-04-30",
            "teams": [{"club": club, "label": f"JU8-{club}", "age_group": "JU8"} for club in CLUBS],
            "age_groups": ["JU8"],
            "round_length_minutes": {"JU8": 10},
            "ice_time_minutes": {"JU8": 120},
        },
        status=StageStatus.DONE,
    )
    state.write_stage(
        StageName.SCRAPING,
        {
            "sources": [
                {"name": "Alfa", "event_count": 3, "url": "https://example.invalid/alfa.ics"},
                {"name": "Bravo", "event_count": 2, "url": "https://example.invalid/bravo.ics"},
            ],
            "blocked": ["Charlie"],
            "confidence": {"verdict": "OK", "overall_assessment": "alt ok"},
            "events_by_club": {},
            "club_calendar_status": {club: "known" for club in CLUBS},
        },
        status=StageStatus.DONE,
    )
    ScrapedDataCache(str(work_dir)).write(
        {
            "_meta": {"updated_at": "2026-09-18T15:56:09", "start_date": "2026-09-01", "end_date": "2027-04-30"},
            "source_count": 2,
            "total_events": 5,
            "sources": {
                "Alfa": {
                    "name": "Alfa",
                    "url": "https://example.invalid/alfa.ics",
                    "scrape_timestamp": "2026-09-18T15:56:09",
                    "event_count": 3,
                    "blocked": False,
                    "events": [
                        {"date": "05.10.2026", "name": "Trening", "datetime": "2026-10-05T18:00:00", "duration_hours": 1.0},
                    ],
                },
                "Bravo": {
                    "name": "Bravo",
                    "url": "https://example.invalid/bravo.ics",
                    "scrape_timestamp": "2026-09-18T15:56:09",
                    "event_count": 2,
                    "blocked": False,
                    "events": [
                        {"date": "06.10.2026", "name": "Kamp", "datetime": "2026-10-06T18:00:00", "duration_hours": 1.0},
                    ],
                },
            },
        }
    )
    state.write_stage(StageName.PLANNING, {"plan": _plan()}, status=StageStatus.DONE)
    return state


def test_stage4_export_stores_public_export_context(tmp_path):
    state = _seed_workspace(tmp_path)
    result = run_export(
        {"plan": _plan()}, state, export_dir=str(tmp_path / "export"), timestamped_export=False
    )

    context = result["public_export_context"]
    assert context["scrape"]["source_count"] == 2
    assert context["scrape"]["total_events"] == 5
    assert context["scrape"]["blocked"] == ["Charlie"]
    assert context["input"]["file_name"] == "input.xlsx"
    assert len(context["input"]["teams"]) == 4
    assert context["activities"]["activities"][0]["title"] == "Spillerutvikling JU8"
    assert result["public_export_context_fingerprint"] == fingerprint_public_export_context(context)


def test_promotion_persists_public_export_context_and_fingerprint(tmp_path):
    state = _seed_workspace(tmp_path)
    result = run_export(
        {"plan": _plan()}, state, export_dir=str(tmp_path / "export"), timestamped_export=False
    )
    schedule, _decisions = promote_from_stage3(
        work_dir=state.work_dir, root=tmp_path / "season", actor="tester"
    )

    assert schedule["promoted_from"]["public_export_context_fingerprint"] == result[
        "public_export_context_fingerprint"
    ]
    stored = resolve_promoted_public_export_context(
        schedule, load_export_context("2026-2027", root=tmp_path / "season")
    )
    assert stored["scrape"]["total_events"] == 5

    export_context_path = tmp_path / "season" / "2026-2027" / "export_context.json"
    assert export_context_path.exists()
    assert json.loads(export_context_path.read_text(encoding="utf-8"))["scrape"]["source_count"] == 2


def test_tampered_public_export_context_is_refused(tmp_path):
    state = _seed_workspace(tmp_path)
    run_export({"plan": _plan()}, state, export_dir=str(tmp_path / "export"), timestamped_export=False)
    schedule, _decisions = promote_from_stage3(
        work_dir=state.work_dir, root=tmp_path / "season", actor="tester"
    )

    with pytest.raises(PublicExportContextError):
        verify_public_export_context(
            {"scrape": {"source_count": 99}},
            expected_fingerprint=schedule["promoted_from"]["public_export_context_fingerprint"],
        )


def test_canonical_export_refuses_tampered_export_context_file(tmp_path, capsys):
    from tournament_scheduler.cli.rvv_cli import main

    state = _seed_workspace(tmp_path)
    run_export({"plan": _plan()}, state, export_dir=str(tmp_path / "export"), timestamped_export=False)
    assert main(["season", "promote", "--work-dir", str(state.work_dir), "--root", str(tmp_path / "season")]) == 0
    capsys.readouterr()

    context_path = tmp_path / "season" / "2026-2027" / "export_context.json"
    tampered = json.loads(context_path.read_text(encoding="utf-8"))
    tampered["scrape"]["source_count"] = 99
    context_path.write_text(json.dumps(tampered), encoding="utf-8")

    rc = main([
        "season",
        "export",
        "--season",
        "2026-2027",
        "--work-dir",
        str(state.work_dir),
        "--root",
        str(tmp_path / "season"),
        "--export-dir",
        str(tmp_path / "canonical-export"),
        "--flat",
    ])
    assert rc == 1
    assert "does not match its recorded fingerprint" in capsys.readouterr().out


def test_canonical_export_after_workspace_deleted_retains_source_and_companions(tmp_path, capsys):
    from tournament_scheduler.cli.rvv_cli import main

    state = _seed_workspace(tmp_path)
    run_export({"plan": _plan()}, state, export_dir=str(tmp_path / "export"), timestamped_export=False)
    assert main(["season", "promote", "--work-dir", str(state.work_dir), "--root", str(tmp_path / "season")]) == 0
    capsys.readouterr()
    # Prove canonical export is self-contained: the entire live workspace
    # (Stage 1/2/3/4 + scrape cache) is gone before exporting.
    shutil.rmtree(state.work_dir)

    rc = main([
        "season",
        "export",
        "--season",
        "2026-2027",
        "--work-dir",
        str(state.work_dir),
        "--root",
        str(tmp_path / "season"),
        "--export-dir",
        str(tmp_path / "canonical-export"),
        "--flat",
        "--json",
    ])
    assert rc == 0
    captured = capsys.readouterr().out
    result = json.loads(captured[captured.index("{") : captured.rindex("}") + 1])

    assert result["verify_result"]["ok"] is True
    files = result["output_files"]
    assert "calendars_html" in files
    assert "input_html" in files
    assert "activities_json" in files

    season_html = Path(files["html"]).read_text(encoding="utf-8")
    assert "2 kilder" in season_html
    assert "5 hendelser" in season_html
    assert "0 kilder" not in season_html
    assert 'href="calendars.html"' in season_html
    assert 'href="input.html"' in season_html
    # Plan-local counts stay in the shared navbar status.
    assert "1 turneringer" in season_html
    assert "6 kamper" in season_html
    assert "4 lag" in season_html

    report_html = Path(files["html_report"]).read_text(encoding="utf-8")
    assert 'href="calendars.html"' in report_html
    assert 'href="input.html"' in report_html

    calendars_html = Path(files["calendars_html"]).read_text(encoding="utf-8")
    assert "Alfa" in calendars_html
    input_html = Path(files["input_html"]).read_text(encoding="utf-8")
    assert "JU8-Alfa" in input_html
    assert Path(files["activities_json"]).exists()


def test_canonical_export_uses_publication_history_not_requested_export_dir(tmp_path, capsys):
    from tournament_scheduler.cli.rvv_cli import main

    state = _seed_workspace(tmp_path)
    initial = run_export({"plan": _plan()}, state, export_dir=str(tmp_path / "export"))
    published_dir = Path(initial["export_dir"])
    manifest_path = published_dir / "export_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest.update(
        {
            "lifecycle_status": "published",
            "canonical_season": "2026-2027",
            "canonical_revision": "published-rev",
            "published_at": "2026-09-21T09:14:53+00:00",
        }
    )
    # Simulate the real committed legacy baseline: published before manifests
    # had a schedule_projection and without the HTML file, while the Spond
    # workbook can still be bound back to stable canonical ids.
    manifest.pop("schedule_projection", None)
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    (published_dir / "season_plan.html").unlink()

    assert main(["season", "promote", "--work-dir", str(state.work_dir), "--root", str(tmp_path / "season")]) == 0
    capsys.readouterr()

    schedule_path = tmp_path / "season" / "2026-2027" / "schedule.json"
    schedule = json.loads(schedule_path.read_text(encoding="utf-8"))
    # The legacy publication predates schedule_projection, so its stable-id
    # baseline must be recovered from the canonical plan at the recorded
    # publication revision. Record that publication-time snapshot durably.
    revisions_dir = schedule_path.parent.parent / "revisions" / "2026-2027"
    revisions_dir.mkdir(parents=True)
    (revisions_dir / "published-rev.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "season": "2026-2027",
                "revision": "published-rev",
                "plan": schedule["plan"],
                # The versioned operational projection backfills the
                # publication-time occupied interval from this snapshot's
                # bound verification context, exactly like the durable
                # canonical schedule does.
                "verification_context": schedule["verification_context"],
            }
        ),
        encoding="utf-8",
    )
    tournament = schedule["plan"]["tournaments"][0]
    tournament["date"] = "2026-09-13"
    schedule["revision"] = "current-rev"
    problem = schedule["verification_context"]["problem"]
    problem["club_calendar_status"] = {club: "known" for club in CLUBS}
    problem["club_busy_intervals"] = {
        "Alfa": [
            {
                "date": "2026-09-13",
                "start": "10:30",
                "end": "12:00",
                "availability": "fixed_busy",
                "kind": "external",
                "calendar_event": "Later booking",
            }
        ]
    }
    schedule_path.write_text(json.dumps(schedule), encoding="utf-8")

    other_export_dir = tmp_path / "fresh-empty-output"
    rc = main([
        "season",
        "export",
        "--season",
        "2026-2027",
        "--work-dir",
        str(state.work_dir),
        "--root",
        str(tmp_path / "season"),
        "--export-dir",
        str(other_export_dir),
        "--flat",
        "--json",
    ])
    assert rc == 0
    captured = capsys.readouterr().out
    result = json.loads(captured[captured.index("{") : captured.rindex("}") + 1])

    guard = result["export_projection_guard"]
    assert guard["published_to_canonical_delta"] is not None
    assert guard["published_to_canonical_delta"]["placement_changes"][0]["tournament_id"] == "ju8-alfa-20260912"
    assert guard["canonical_delta"]["changed"] is False
    assert result.get("placement_normalization") in (None, {})
    assert result["reviewed_plan"]["tournaments"][0]["date"] == "2026-09-13"


def test_canonical_export_fails_when_legacy_published_projection_is_unrecoverable(tmp_path, capsys):
    from tournament_scheduler.cli.rvv_cli import main

    state = _seed_workspace(tmp_path)
    run_export({"plan": _plan()}, state, export_dir=str(tmp_path / "export"), timestamped_export=False)
    history_dir = tmp_path / "export" / "published-without-artifacts"
    history_dir.mkdir(parents=True)
    (history_dir / "export_manifest.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "export_id": "published-without-artifacts",
                "generated_at": "2026-09-21T09:08:01+00:00",
                "canonical_season": "2026-2027",
                "canonical_revision": "published-rev",
                "lifecycle_status": "published",
            }
        ),
        encoding="utf-8",
    )
    assert main(["season", "promote", "--work-dir", str(state.work_dir), "--root", str(tmp_path / "season")]) == 0
    capsys.readouterr()

    rc = main([
        "season",
        "export",
        "--season",
        "2026-2027",
        "--work-dir",
        str(state.work_dir),
        "--root",
        str(tmp_path / "season"),
        "--export-dir",
        str(tmp_path / "fresh-empty-output"),
        "--flat",
    ])
    assert rc == 1
    assert "could not be resolved" in capsys.readouterr().out


def test_canonical_export_does_not_read_mutated_workspace_scrape_state(tmp_path, capsys):
    from tournament_scheduler.cli.rvv_cli import main

    state = _seed_workspace(tmp_path)
    run_export({"plan": _plan()}, state, export_dir=str(tmp_path / "export"), timestamped_export=False)
    assert main(["season", "promote", "--work-dir", str(state.work_dir), "--root", str(tmp_path / "season")]) == 0
    capsys.readouterr()

    # A later scrape writes completely different counts into the live cache.
    ScrapedDataCache(str(state.work_dir)).write(
        {"_meta": {"updated_at": "2027-01-01T00:00:00"}, "source_count": 9, "total_events": 999, "sources": {}}
    )

    rc = main([
        "season",
        "export",
        "--season",
        "2026-2027",
        "--work-dir",
        str(state.work_dir),
        "--root",
        str(tmp_path / "season"),
        "--export-dir",
        str(tmp_path / "canonical-export"),
        "--flat",
        "--json",
    ])
    assert rc == 0
    captured = capsys.readouterr().out
    result = json.loads(captured[captured.index("{") : captured.rindex("}") + 1])
    season_html = Path(result["output_files"]["html"]).read_text(encoding="utf-8")
    assert "2 kilder" in season_html
    assert "999 hendelser" not in season_html
