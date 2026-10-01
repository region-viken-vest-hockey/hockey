"""Tests for tournament_scheduler.pipeline.pages_bundle (issue #18)."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from tournament_scheduler.html.data_computation import build_export_links_html
from tournament_scheduler.pipeline.pages_bundle import (
    DEFAULT_ALLOWED_DIRECTORIES,
    DEFAULT_ALLOWED_FILENAMES,
    build_public_bundle,
)


def _export_dir(tmp_path: Path) -> Path:
    d = tmp_path / "export"
    d.mkdir()
    return d


def _write_xlsx(
    path: Path,
    *,
    extra_row: list[str] | None = None,
    hidden_sheet: bool = False,
    formula: bool = False,
) -> None:
    """Write a real, minimal public-safe XLSX workbook for bundle tests."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sesongplan"
    sheet.append(["Dato", "Aktivitet"])
    sheet.append(["10.10.2026", "U10 Turnering \u2014 Jar Isforum"])
    if extra_row is not None:
        sheet.append(extra_row)
    if formula:
        workbook.create_sheet("Formel")["A1"] = "=SUM(1,2)"
    if hidden_sheet:
        hidden = workbook.create_sheet("Skjult")
        hidden["A1"] = "intern"
        hidden.sheet_state = "hidden"
    workbook.save(path)


class TestDefaultAllowlist:
    def test_public_download_exports_are_included_by_default(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text(
            '<a href="season_plan.xlsx" class="export-link-btn">Excel</a>'
            '<a href="season_plan.csv" class="export-link-btn">CSV</a>'
            '<a href="season_plan_overview.csv" class="export-link-btn">CSV overview</a>'
            '<a href="season_plan.ics" class="export-link-btn">iCal</a>',
            encoding="utf-8",
        )
        (export_dir / "season_plan.ics").write_text("BEGIN:VCALENDAR\nEND:VCALENDAR\n", encoding="utf-8")
        _write_xlsx(export_dir / "season_plan.xlsx")
        (export_dir / "season_plan.csv").write_text("date,arena\n", encoding="utf-8")
        (export_dir / "season_plan_overview.csv").write_text("date,arena\n", encoding="utf-8")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        included = set(json.loads(Path((tmp_path / "pages_privacy_report.json")).read_text())["included_files"])
        assert included == {
            "season_plan.html",
            "season_plan.ics",
            "season_plan.xlsx",
            "season_plan.csv",
            "season_plan_overview.csv",
        }
        assert (tmp_path / "public" / "season_plan.xlsx").exists()
        assert (tmp_path / "public" / "season_plan.csv").exists()
        assert (tmp_path / "public" / "season_plan_overview.csv").exists()
        content = (tmp_path / "public" / "season_plan.html").read_text(encoding="utf-8")
        assert 'href="season_plan.xlsx"' in content
        assert 'href="season_plan.csv"' in content
        assert 'href="season_plan_overview.csv"' in content
        assert 'aria-disabled="true"' not in content

    def test_input_html_is_included_by_default(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text("<h1>Plan</h1>", encoding="utf-8")
        (export_dir / "input.html").write_text("<h1>Påmeldte lag</h1>", encoding="utf-8")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        assert "input.html" in DEFAULT_ALLOWED_FILENAMES
        included = set(json.loads(Path((tmp_path / "pages_privacy_report.json")).read_text())["included_files"])
        assert "input.html" in included
        assert (tmp_path / "public" / "input.html").exists()

    def test_season_changes_html_is_included_by_default(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text(
            '<a href="season_changes.html">Forespurte endringer</a>', encoding="utf-8"
        )
        (export_dir / "season_changes.html").write_text("<h1>Forespurte endringer</h1>", encoding="utf-8")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        assert "season_changes.html" in DEFAULT_ALLOWED_FILENAMES
        included = set(json.loads(Path((tmp_path / "pages_privacy_report.json")).read_text())["included_files"])
        assert "season_changes.html" in included
        assert (tmp_path / "public" / "season_changes.html").exists()

    def test_registered_team_artifacts_are_included_without_private_validation_report(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        registered_dir = export_dir / "registered-teams"
        registered_dir.mkdir()
        (registered_dir / "pameldte-lag.html").write_text(
            '<h1>Påmeldte lag</h1><a href="pameldte-lag.json">JSON</a>', encoding="utf-8"
        )
        (registered_dir / "pameldte-lag.json").write_text('{"total_teams": 1}', encoding="utf-8")
        (registered_dir / "validation-report.json").write_text(
            '{"source_sha256": "abc", "source_file": "SharePoint.csv"}', encoding="utf-8"
        )

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        assert "registered-teams" in DEFAULT_ALLOWED_DIRECTORIES
        assert (tmp_path / "public" / "registered-teams" / "pameldte-lag.html").exists()
        assert (tmp_path / "public" / "registered-teams" / "pameldte-lag.json").exists()
        assert not (tmp_path / "public" / "registered-teams" / "validation-report.json").exists()
        report = json.loads((tmp_path / "pages_privacy_report.json").read_text())
        included = set(report["included_files"])
        assert "registered-teams/pameldte-lag.html" in included
        assert "registered-teams/pameldte-lag.json" in included
        assert any(
            e["file"] == "registered-teams/validation-report.json" and "private validation" in e["reason"]
            for e in report["excluded_files"]
        )

    def test_activity_artifacts_are_included_by_default(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text('<a href="activities/">Aktiviteter</a>', encoding="utf-8")
        (export_dir / "activities.json").write_text('{"activities": []}', encoding="utf-8")
        activities_dir = export_dir / "activities"
        activities_dir.mkdir()
        (activities_dir / "index.html").write_text('<script src="../activities.json"></script>', encoding="utf-8")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        assert "activities.json" in DEFAULT_ALLOWED_FILENAMES
        assert "activities" in DEFAULT_ALLOWED_DIRECTORIES
        assert (tmp_path / "public" / "activities.json").exists()
        assert (tmp_path / "public" / "activities" / "index.html").exists()
        included = set(json.loads(Path((tmp_path / "pages_privacy_report.json")).read_text())["included_files"])
        assert "activities.json" in included
        assert "activities/index.html" in included

    def test_unknown_file_inside_activity_directory_is_excluded(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        activities_dir = export_dir / "activities"
        activities_dir.mkdir()
        (activities_dir / "index.html").write_text("<h1>ok</h1>", encoding="utf-8")
        (activities_dir / "debug.exe").write_bytes(b"MZ")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        assert (tmp_path / "public" / "activities" / "index.html").exists()
        assert not (tmp_path / "public" / "activities" / "debug.exe").exists()
        report = json.loads((tmp_path / "pages_privacy_report.json").read_text())
        assert any(e["file"] == "activities/debug.exe" and "unknown" in e["reason"] for e in report["excluded_files"])

    def test_never_copies_the_whole_export_directory(self, tmp_path):
        """Only allowlisted artifacts are copied; review packets stay private."""
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text("<h1>Plan</h1>", encoding="utf-8")
        _write_xlsx(export_dir / "season_plan.xlsx")
        _write_xlsx(export_dir / "season_plan_spond.xlsx")
        _write_xlsx(export_dir / "season_plan_spond_games.xlsx")
        (export_dir / "season_plan.csv").write_text("club,date\n", encoding="utf-8")
        (export_dir / "season_plan_overview.csv").write_text("club,date\n", encoding="utf-8")
        review_dir = export_dir / "review_packets"
        review_dir.mkdir()
        (review_dir / "club_a.xlsx").write_bytes(b"secret roster data")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        public_files = {p.name for p in (tmp_path / "public").iterdir()}
        assert public_files == {
            "season_plan.html",
            "season_plan.xlsx",
            "season_plan_spond.xlsx",
            "season_plan_spond_games.xlsx",
            "season_plan.csv",
            "season_plan_overview.csv",
        }
        assert not (tmp_path / "public" / "review_packets").exists()

    def test_unknown_extension_is_excluded_even_if_allowlisted_by_name(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.exe").write_bytes(b"MZ")

        result = build_public_bundle(
            str(export_dir), str(tmp_path / "public"), allowed_filenames={"season_plan.exe"}
        )

        assert result.status == "ok"
        assert not (tmp_path / "public" / "season_plan.exe").exists()
        report = json.loads((tmp_path / "pages_privacy_report.json").read_text())
        assert any("unknown" in e["reason"] for e in report["excluded_files"])


class TestSecretDetectionBlocksPublication:
    def test_aws_key_blocks_and_leaves_no_bundle(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text(
            "<p>key=AKIAABCDEFGHIJKLMNOP</p>", encoding="utf-8"
        )

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "blocked"
        assert result.requires_human is True
        assert not (tmp_path / "public").exists()

    def test_bearer_url_blocks_publication(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text(
            '<a href="https://example.com/feed?access_token=abcdef123456">link</a>', encoding="utf-8"
        )

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "blocked"

    def test_allow_findings_overrides_a_specific_false_positive(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text(
            "<p>password=placeholder-not-a-real-secret</p>", encoding="utf-8"
        )

        blocked = build_public_bundle(str(export_dir), str(tmp_path / "public"))
        assert blocked.status == "blocked"

        allowed = build_public_bundle(
            str(export_dir),
            str(tmp_path / "public"),
            allow_findings={"placeholder-not-a-real-secret"},
        )
        assert allowed.status == "ok"
        assert (tmp_path / "public" / "season_plan.html").exists()


class TestRedaction:
    def test_local_filesystem_path_is_redacted_not_blocking(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text(
            "<p>Generated from /Users/alice/hockey/input.xlsx</p>", encoding="utf-8"
        )

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        content = (tmp_path / "public" / "season_plan.html").read_text(encoding="utf-8")
        assert "/Users/alice" not in content
        assert "[redacted]" in content

    def test_contact_email_is_redacted(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text(
            "<p>Contact: organizer@example.com</p>", encoding="utf-8"
        )

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        content = (tmp_path / "public" / "season_plan.html").read_text(encoding="utf-8")
        assert "organizer@example.com" not in content

    def test_labeled_phone_number_is_redacted(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text("<p>Tlf: 123 45 678</p>", encoding="utf-8")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        content = (tmp_path / "public" / "season_plan.html").read_text(encoding="utf-8")
        assert "123 45 678" not in content

    def test_ordinary_dates_and_numbers_are_not_redacted(self, tmp_path):
        """Unlabeled digit sequences (dates, scores) must survive untouched."""
        export_dir = _export_dir(tmp_path)
        original = "<p>2026-03-05 14:00 — Jar 3 - 2 Bekkelaget, hall 12345</p>"
        (export_dir / "season_plan.html").write_text(original, encoding="utf-8")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        content = (tmp_path / "public" / "season_plan.html").read_text(encoding="utf-8")
        assert content == original


class TestAssetRewriting:
    def test_root_absolute_links_are_rewritten_relative(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan.html").write_text(
            '<link rel="stylesheet" href="/styles.css"><img src="/logo.png">', encoding="utf-8"
        )

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        content = (tmp_path / "public" / "season_plan.html").read_text(encoding="utf-8")
        assert 'href="styles.css"' in content
        assert 'src="logo.png"' in content

    def test_external_and_protocol_relative_links_are_untouched(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        original = '<a href="https://example.com/page">x</a><script src="//cdn.example.com/a.js">'
        (export_dir / "season_plan.html").write_text(original, encoding="utf-8")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        content = (tmp_path / "public" / "season_plan.html").read_text(encoding="utf-8")
        assert content == original


class TestExcludedFileLinkRewriting:
    def test_excluded_file_links_are_removed_and_reported(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        review_packets = export_dir / "review_packets"
        review_packets.mkdir()
        (review_packets / "club_a.xlsx").write_bytes(b"review data")
        _write_xlsx(export_dir / "season_plan.xlsx")
        (export_dir / "season_plan.csv").write_text("club,date\n", encoding="utf-8")
        (export_dir / "season_plan_overview.csv").write_text("club,date\n", encoding="utf-8")
        (export_dir / "season_plan_team_counts.csv").write_text("club,count\n", encoding="utf-8")
        (export_dir / "season_plan.html").write_text(
            (
                '<a href="season_plan.xlsx" class="export-link-btn">Excel</a>'
                '<a href="season_plan.csv" class="export-link-btn">CSV</a>'
                '<a href="season_plan_overview.csv" class="export-link-btn">CSV overview</a>'
                '<a href="review_packets/club_a.xlsx">Review packet</a>'
                '<a href="calendars.html">Calendar</a>'
                '<a href="https://example.com/page">External</a>'
                '<img src="review_packets/club_a.xlsx">'
                '<img src="season_plan_team_counts.csv">'
            ),
            encoding="utf-8",
        )

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        content = (tmp_path / "public" / "season_plan.html").read_text(encoding="utf-8")
        assert '<a href="season_plan.xlsx" class="export-link-btn">Excel</a>' in content
        assert '<a href="season_plan.csv" class="export-link-btn">CSV</a>' in content
        assert '<a href="season_plan_overview.csv" class="export-link-btn">CSV overview</a>' in content
        assert 'aria-disabled="true"' in content
        assert 'data-excluded-href="review_packets/club_a.xlsx"' in content
        assert 'data-excluded-src="review_packets/club_a.xlsx"' in content
        assert 'data-excluded-src="season_plan_team_counts.csv"' in content
        assert 'href="calendars.html"' in content
        assert 'href="https://example.com/page"' in content

        report = json.loads((tmp_path / "pages_privacy_report.json").read_text())
        rewrites = report["rewritten_links"]
        assert any(
            entry["target"] == "review_packets/club_a.xlsx"
            and entry["attribute"] == "href"
            and entry["action"] == "disabled"
            for entry in rewrites
        )
        assert any(
            entry["target"] == "review_packets/club_a.xlsx"
            and entry["scope"] == "excluded_directory"
            for entry in rewrites
        )
        assert any(
            entry["target"] == "season_plan_team_counts.csv"
            and entry["attribute"] == "src"
            and entry["action"] == "removed"
            for entry in rewrites
        )


class TestMissingExportDir:
    def test_missing_export_dir_returns_failed(self, tmp_path):
        result = build_public_bundle(str(tmp_path / "nope"), str(tmp_path / "public"))
        assert result.status == "failed"


class TestDefaultAllowlistConstant:
    def test_default_allowlist_includes_public_download_formats(self):
        assert "season_plan.xlsx" in DEFAULT_ALLOWED_FILENAMES
        assert "season_plan.csv" in DEFAULT_ALLOWED_FILENAMES
        assert "season_plan_overview.csv" in DEFAULT_ALLOWED_FILENAMES
        assert "season_plan_spond.xlsx" in DEFAULT_ALLOWED_FILENAMES
        assert "season_plan_spond_games.xlsx" in DEFAULT_ALLOWED_FILENAMES


def test_spond_download_buttons_render_when_present():
    links = build_export_links_html(
        {
            "excel": "/tmp/export/season_plan.xlsx",
            "spond": "/tmp/export/season_plan_spond.xlsx",
            "spond_games": "/tmp/export/season_plan_spond_games.xlsx",
        }
    )

    assert 'href="season_plan_spond.xlsx"' in links
    assert 'href="season_plan_spond_games.xlsx"' in links
    assert "data-excluded" not in links


def test_spond_download_buttons_omitted_when_absent():
    links = build_export_links_html({"excel": "/tmp/export/season_plan.xlsx"})

    assert "spond" not in links


class TestSpondPublicBundle:
    def test_spond_workbooks_and_links_survive_sanitization(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        _write_xlsx(export_dir / "season_plan_spond.xlsx")
        _write_xlsx(export_dir / "season_plan_spond_games.xlsx")
        (export_dir / "season_plan.html").write_text(
            '<a href="season_plan_spond.xlsx" class="export-link-btn" download>Spond sesongplan</a>'
            '<a href="season_plan_spond_games.xlsx" class="export-link-btn" download>Spond kampoppsett</a>',
            encoding="utf-8",
        )

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        assert (tmp_path / "public" / "season_plan_spond.xlsx").exists()
        assert (tmp_path / "public" / "season_plan_spond_games.xlsx").exists()
        content = (tmp_path / "public" / "season_plan.html").read_text(encoding="utf-8")
        assert 'href="season_plan_spond.xlsx"' in content
        assert 'href="season_plan_spond_games.xlsx"' in content
        assert "aria-disabled" not in content
        assert "data-excluded-href" not in content

    def test_sensitive_cell_value_blocks_the_bundle(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        _write_xlsx(export_dir / "season_plan_spond.xlsx", extra_row=["x", "organizer@example.com"])

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "blocked"
        assert result.requires_human is True
        assert not (tmp_path / "public").exists()

    def test_hidden_sheet_blocks_the_bundle(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        _write_xlsx(export_dir / "season_plan_spond.xlsx", hidden_sheet=True)

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "blocked"
        assert any("hidden" in problem for problem in result.problems)

    def test_formula_blocks_the_bundle(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        _write_xlsx(export_dir / "season_plan_spond_games.xlsx", formula=True)

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "blocked"
        assert not (tmp_path / "public").exists()

    def test_unreadable_workbook_blocks_the_bundle(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        (export_dir / "season_plan_spond.xlsx").write_bytes(b"not a real workbook")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "blocked"
        assert not (tmp_path / "public").exists()


def _workbook_with_comment(path: Path) -> None:
    from openpyxl.comments import Comment

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sesongplan"
    sheet["A1"] = "approved"
    sheet["A1"].comment = Comment("private roster note: coach@example.com", "author")
    workbook.save(path)


def _workbook_with_hyperlink(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sesongplan"
    sheet["A1"] = "approved"
    sheet["A1"].hyperlink = "https://internal.example/booking-evidence"
    workbook.save(path)


def _inject_zip_part(path: Path, name: str, data: bytes) -> None:
    import zipfile

    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(tmp, "w") as destination:
        for item in source.infolist():
            destination.writestr(item, source.read(item.filename))
        destination.writestr(name, data)
    tmp.replace(path)


class TestXlsxPackageValidation:
    """Issue #509 follow-up: a binary workbook must not be copied on trust.

    Text can hide in comments, hyperlink targets and non-cell package parts, so
    the whole package is validated and any unexpected part blocks the bundle.
    """

    def test_sensitive_comment_in_season_workbook_blocks_the_bundle(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        _workbook_with_comment(export_dir / "season_plan.xlsx")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "blocked"
        assert result.requires_human is True
        assert not (tmp_path / "public").exists()
        assert any("comment" in problem for problem in result.problems)

    def test_hyperlink_in_season_workbook_blocks_the_bundle(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        _workbook_with_hyperlink(export_dir / "season_plan.xlsx")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "blocked"
        assert not (tmp_path / "public").exists()

    def test_unexpected_package_part_blocks_the_bundle(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        workbook_path = export_dir / "season_plan.xlsx"
        _write_xlsx(workbook_path)
        _inject_zip_part(workbook_path, "customXml/item1.xml", b"<private>note</private>")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "blocked"
        assert not (tmp_path / "public").exists()
        assert any("custom_xml" in problem or "unexpected" in problem for problem in result.problems)

    def test_clean_workbooks_still_pass(self, tmp_path):
        export_dir = _export_dir(tmp_path)
        _write_xlsx(export_dir / "season_plan.xlsx")
        _write_xlsx(export_dir / "season_plan_spond.xlsx")
        _write_xlsx(export_dir / "season_plan_spond_games.xlsx")

        result = build_public_bundle(str(export_dir), str(tmp_path / "public"))

        assert result.status == "ok"
        assert (tmp_path / "public" / "season_plan.xlsx").exists()
        assert (tmp_path / "public" / "season_plan_spond.xlsx").exists()
        assert (tmp_path / "public" / "season_plan_spond_games.xlsx").exists()
