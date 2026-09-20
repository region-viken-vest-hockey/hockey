"""Full multi-process Stage 3 lifecycle regression (hermetic, no network).

This is the independent process-boundary safety net for the interactive
``rvv-miniputt run --interactive`` protocol. Every transition below happens in
a *separate* OS process, so it exercises exactly what a thin harness adapter
sees across checkpoints rather than a single in-process unit path:

    Stage 1 -> Stage 2 -> shared-host sub-decision -> Stage 3 baseline candidate
      -> optimize_plan (search) -> optimize_plan (different search)
      -> apply_candidate (finalize exact revision) -> Stage 4 export

It asserts the canonical lifecycle facts the CLI must preserve across those
boundaries: run id ownership, exact candidate revision/fingerprint lineage,
stale/unknown candidate rejection, wrong-resume-ownership rejection, and the
persisted Stage 3 session exposing the finalized identity Stage 4 consumes.

Marked ``integration`` (subprocess/multi-stage, per the pytest convention in
``pyproject.toml``) so it runs in ``scripts/check full`` but not the quick
suite. It is hermetic: a tiny synthetic workbook, zero calendar sources, and
an isolated canonical-season root, so no external service is ever contacted.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]

pytestmark = pytest.mark.integration


def _isolated_env() -> dict[str, str]:
    """Return a subprocess environment with no ambient operator/harness state.

    The canonical-season root is pinned to a non-existent path so the
    lifecycle cannot silently plan the developer's promoted ``season/`` state
    instead of the synthetic workbook, and every harness/LLM-judge variable is
    dropped so each decision must come from the explicit action we submit.
    """
    env = dict(os.environ)
    env["RVV_CANONICAL_SEASON_ROOT"] = "/nonexistent/rvv-lifecycle-season"
    for name in ("RVV_HARNESS", "PI_SESSION_ID", "CLAUDE_CODE_SESSION_ID", "RVV_JUDGE_BACKEND"):
        env.pop(name, None)
    return env


def _write_minimal_workbook(path: Path) -> None:
    # Mirrors tests/test_cli_smoke.py: the season window must stay ahead of the
    # wall clock because Stage 3 clamps planning to a future effective start.
    start = date.today() + timedelta(days=30)
    end = start + timedelta(days=105)

    wb = openpyxl.Workbook()
    settings = wb.active
    settings.title = "Innstillinger"
    settings.append(["felt", "verdi"])
    settings.append(["start_date", start.isoformat()])
    settings.append(["end_date", end.isoformat()])

    age_groups = wb.create_sheet("Aldersgrupper")
    age_groups.append(
        [
            "age_group",
            "parallel_games",
            "round_length_minutes",
            "ice_time_minutes",
            "deltakelser_per_lag_før_jul",
            "deltakelser_per_lag_etter_jul",
        ]
    )
    age_groups.append(["U10", 2, 15, 30, 2, 2])

    teams = wb.create_sheet("Lag")
    teams.append(["club", "label", "age_group"])
    for club, label in (
        ("Kongsberg", "Kongsberg U10A"),
        ("Skien", "Skien U10A"),
        ("Ringerike", "Ringerike U10A"),
        ("Jar", "Jar U10A"),
        # A joint registration ("A/B") forces the run-scoped shared-host
        # sub-decision the lifecycle must resolve before building the plan.
        ("Kongsberg/Skien", "Kongsberg Skien U10A"),
    ):
        teams.append([club, label, "U10"])

    # Deliberately no "Kilder" sheet: Stage 2 has nothing to scrape, so the
    # whole lifecycle makes zero network calls.
    wb.save(path)


def _run_cli(args: list[str], *, timeout: int = 600) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "tournament_scheduler.cli.rvv_cli", *args],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=timeout,
        env=_isolated_env(),
    )


def _trailing_decision_context(stdout: str) -> dict:
    """Parse the DecisionContext JSON emitted as the last stdout block."""
    marker = stdout.rfind("\n{\n")
    assert marker != -1, f"no DecisionContext JSON found in stdout:\n{stdout}"
    return json.loads(stdout[marker:])


def _session_view(work_dir: Path) -> dict:
    """Read the canonical Stage 3 session facade, never the raw side file."""
    result = _run_cli(["stage3", "session", "--work-dir", str(work_dir), "--json"])
    assert result.returncode == 0, result.stdout + result.stderr
    return json.loads(result.stdout)


class TestInteractiveStage3LifecycleAcrossProcesses:
    def _interactive_run(
        self,
        base: list[str],
        *,
        resume: int,
        action: dict | None = None,
        expect_error: bool = False,
    ) -> subprocess.CompletedProcess[str]:
        argv = [*base, "--resume-from", str(resume)]
        if action is not None:
            argv += ["--decision-action", json.dumps(action)]
        result = _run_cli(argv)
        if expect_error:
            assert result.returncode == 1, result.stdout + result.stderr
            assert "Traceback" not in result.stderr
        else:
            assert result.returncode == 2, result.stdout + result.stderr
            assert "Traceback" not in result.stderr
        return result

    def test_exact_revision_lineage_across_process_boundaries(self, tmp_path):
        input_path = tmp_path / "input.xlsx"
        _write_minimal_workbook(input_path)
        work_dir = tmp_path / ".pipeline"
        export_dir = tmp_path / "export"

        base = [
            "run",
            "--interactive",
            "--input",
            str(input_path),
            "--work-dir",
            str(work_dir),
            "--export-dir",
            str(export_dir),
            "--non-strict",
            "--allow-missing-sources",
            "--no-timestamped-export",
        ]

        # --- Stage 1 (own process) ---------------------------------------
        ctx = _trailing_decision_context(self._interactive_run(base, resume=1).stdout)
        assert ctx["capability"] == "config"

        # --- Stage 2 (own process) ---------------------------------------
        ctx = _trailing_decision_context(
            self._interactive_run(base, resume=2, action={"action_id": "proceed"}).stdout
        )
        assert ctx["capability"] == "scraping"

        # --- Stage 3: run-scoped shared-host sub-decision (own process) --
        ctx = _trailing_decision_context(
            self._interactive_run(base, resume=3, action={"action_id": "proceed"}).stdout
        )
        assert ctx["capability"] == "shared_host_assignment"
        chosen_club = ctx["action_parameters"]["assign_shared_host"]["chosen_club"]["enum"][0]

        # Answering the run-scoped decision with the same resume number lets
        # the *same* process continue into Stage 3 and emit the candidate.
        ctx = _trailing_decision_context(
            self._interactive_run(
                base,
                resume=3,
                action={"action_id": "assign_shared_host", "arguments": {"chosen_club": chosen_club}},
            ).stdout
        )
        assert ctx["capability"] == "stage3_interactive"
        # The config/scraping contexts are not yet scoped to a logical run;
        # the Stage 3 context is, and every later emission must keep it.
        run_id = ctx["run_id"]
        assert run_id

        session = _session_view(work_dir)
        assert session["run_id"] == run_id
        assert session["shared_host_decisions"] == 1
        assert session["candidate_revision"] == 1
        fingerprint_rev1 = session["candidate_fingerprint"]
        assert fingerprint_rev1
        assert session["pending_decision"]["resume_from"] == 4

        # --- Search #1: a materially different bounded attempt ------------
        ctx = _trailing_decision_context(
            self._interactive_run(
                base,
                resume=4,
                action={
                    "action_id": "optimize_plan",
                    "arguments": {"iterations": 300, "seed": 11, "move_dates": True},
                    "rationale": "bounded search one",
                },
            ).stdout
        )
        assert ctx["capability"] == "stage3_optimize"
        assert ctx["candidate_ref"] == "stage3_interactive:attempt_2"

        session = _session_view(work_dir)
        assert session["candidate_revision"] == 2
        fingerprint_rev2 = session["candidate_fingerprint"]
        assert fingerprint_rev2 != fingerprint_rev1

        # --- Search #2: different scope, still legal after search #1 ------
        ctx = _trailing_decision_context(
            self._interactive_run(
                base,
                resume=4,
                action={
                    "action_id": "optimize_plan",
                    "arguments": {"iterations": 300, "seed": 12, "move_hosts": True},
                    "rationale": "bounded search two",
                },
            ).stdout
        )
        assert ctx["capability"] == "stage3_optimize"
        assert ctx["candidate_ref"] == "stage3_interactive:attempt_3"

        session = _session_view(work_dir)
        assert session["candidate_revision"] == 3
        fingerprint_rev3 = session["candidate_fingerprint"]
        assert fingerprint_rev3 not in (fingerprint_rev1, fingerprint_rev2)
        assert session["finalized_revision"] is None

        # --- A stale prior-revision transition is rejected precisely ------
        stale = self._interactive_run(
            base,
            resume=4,
            action={
                "action_id": "apply_candidate",
                "arguments": {
                    "candidate_ref": "stage3_interactive:attempt_1",
                    "candidate_revision": 1,
                    "candidate_fingerprint": fingerprint_rev1,
                },
            },
            expect_error=True,
        )
        assert "stale_candidate" in stale.stdout
        # The rejection must not have replaced the current candidate.
        assert _session_view(work_dir)["candidate_fingerprint"] == fingerprint_rev3

        # --- Finalize the exact current revision through the public path --
        retained = session["retained_candidate_refs"]
        assert retained[-1] == "stage3_interactive:attempt_3"
        result = self._interactive_run(
            base,
            resume=4,
            action={
                "action_id": "apply_candidate",
                "arguments": {
                    "candidate_ref": "stage3_interactive:attempt_3",
                    "candidate_revision": 3,
                    "candidate_fingerprint": fingerprint_rev3,
                },
                "rationale": "adopt the reviewed revision",
            },
        )
        ctx = _trailing_decision_context(result.stdout)
        assert ctx["capability"] == "export"

        # --- The finalized identity Stage 4 consumed is exact -------------
        session = _session_view(work_dir)
        assert session["status"] == "finalized"
        assert session["finalized_revision"] == 3
        assert session["finalized_fingerprint"] == fingerprint_rev3
        assert session["candidate_fingerprint"] == fingerprint_rev3
        assert session["pending_decision"] is None

        export_checkpoint = json.loads((work_dir / "stage4_export.json").read_text(encoding="utf-8"))
        assert export_checkpoint["status"] == "done"
        assert export_checkpoint["data"]["errors"] == []

        # --- Coarse timings: solver/search budget stays separate from the
        # deterministic overhead, so a performance regression is attributable.
        manifest = json.loads((work_dir / "run_manifest.json").read_text(encoding="utf-8"))
        assert manifest["run_id"] == run_id
        timings = manifest["timing"]
        for key in (
            "stage3_baseline_seconds",
            "stage3_local_search_seconds",
            "stage3_verification_seconds",
            "stage3_decision_context_seconds",
            "stage4_export_seconds",
        ):
            assert key in timings, f"missing timing {key}: {sorted(timings)}"
        assert timings["stage3_local_search_seconds"] >= 0

    def test_wrong_resume_ownership_is_rejected_across_the_process_boundary(self, tmp_path):
        """A candidate-scoped Stage 3 decision answered with ``--resume-from 3``
        must fail with an ownership error instead of rebuilding Stage 2."""
        input_path = tmp_path / "input.xlsx"
        _write_minimal_workbook(input_path)
        work_dir = tmp_path / ".pipeline"
        export_dir = tmp_path / "export"

        base = [
            "run",
            "--interactive",
            "--input",
            str(input_path),
            "--work-dir",
            str(work_dir),
            "--export-dir",
            str(export_dir),
            "--non-strict",
            "--allow-missing-sources",
            "--no-timestamped-export",
        ]

        self._interactive_run(base, resume=1)
        self._interactive_run(base, resume=2, action={"action_id": "proceed"})
        shared_host = _trailing_decision_context(
            self._interactive_run(base, resume=3, action={"action_id": "proceed"}).stdout
        )
        assert shared_host["capability"] == "shared_host_assignment"
        chosen_club = shared_host["action_parameters"]["assign_shared_host"]["chosen_club"]["enum"][0]
        self._interactive_run(
            base,
            resume=3,
            action={"action_id": "assign_shared_host", "arguments": {"chosen_club": chosen_club}},
        )

        session = _session_view(work_dir)
        fingerprint = session["candidate_fingerprint"]
        assert session["pending_decision"]["resume_from"] == 4

        # Answer with the wrong (run-scoped) resume number.
        result = self._interactive_run(
            base,
            resume=3,
            action={
                "action_id": "apply_candidate",
                "arguments": {
                    "candidate_ref": session["retained_candidate_refs"][-1],
                    "candidate_revision": session["candidate_revision"],
                    "candidate_fingerprint": fingerprint,
                },
            },
            expect_error=True,
        )
        assert "--resume-from 4" in result.stdout
        assert "decision_action_not_available" not in result.stdout
        # The wrong resume must not have mutated the candidate.
        assert _session_view(work_dir)["candidate_fingerprint"] == fingerprint
