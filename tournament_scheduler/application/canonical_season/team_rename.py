"""Canonical team-identity rename mutations."""

from __future__ import annotations

import copy
from typing import Any, Mapping, Sequence

from tournament_scheduler.canonical_baseline import approval_fingerprint
from tournament_scheduler.canonical_state import canonical_state_revision, schedule_fingerprint
from tournament_scheduler.infrastructure.canonical_season_store import CanonicalSeasonSnapshot, SeasonStateError
from tournament_scheduler.pipeline.fingerprints import stable_payload_sha256
from tournament_scheduler.plan_derived_state import reconcile_plan_derived_state
from tournament_scheduler.planning_contract import verify_candidate
from tournament_scheduler.serialization.season_plan import SEASON_PLAN_SCHEMA_VERSION

from .shared import _operator_identity, _now_iso, _resolve_plan_problem

TeamIdentity = tuple[str, str, str]
RenameMapping = tuple[TeamIdentity, TeamIdentity]


def _identity_from_mapping(raw: Mapping[str, Any], *, prefix: str) -> RenameMapping:
    club = str(raw.get("club") or "").strip()
    age_group = str(raw.get("age_group") or "").strip()
    from_label = str(raw.get("from_label") or raw.get("from") or "").strip()
    to_label = str(raw.get("to_label") or raw.get("to") or "").strip()
    if not club or not age_group or not from_label or not to_label:
        raise SeasonStateError(
            f"Invalid {prefix}: expected club, age_group, from_label/from and to_label/to"
        )
    if from_label == to_label:
        raise SeasonStateError(f"Invalid {prefix}: rename would be a no-op for {club}/{age_group}/{from_label}")
    return (club, from_label, age_group), (club, to_label, age_group)


def _normalize_mappings(mappings: Sequence[Mapping[str, Any]]) -> list[RenameMapping]:
    if not mappings:
        raise SeasonStateError("At least one team rename mapping is required")
    normalized = [_identity_from_mapping(raw, prefix=f"rename mapping #{index + 1}") for index, raw in enumerate(mappings)]
    sources = [source for source, _target in normalized]
    targets = [target for _source, target in normalized]
    if len(set(sources)) != len(sources):
        raise SeasonStateError("Refusing team rename: duplicate source identities are ambiguous")
    if len(set(targets)) != len(targets):
        raise SeasonStateError("Refusing team rename: duplicate target identities would collide")
    source_set = set(sources)
    for source, target in normalized:
        if target in source_set:
            raise SeasonStateError(
                "Refusing team rename: target identities may not be another source in the same atomic rename"
            )
    return normalized


def _team_identity(team: Mapping[str, Any], fallback_age_group: str = "") -> TeamIdentity:
    return (
        str(team.get("club") or ""),
        str(team.get("label") or ""),
        str(team.get("age_group") or fallback_age_group),
    )


def _team_ref(identity: TeamIdentity) -> dict[str, str]:
    return {"club": identity[0], "label": identity[1], "age_group": identity[2]}


def _assert_registered_pool_allows_rename(problem: Mapping[str, Any], mappings: Sequence[RenameMapping]) -> None:
    teams = [team for team in problem.get("teams", []) or [] if isinstance(team, Mapping)]
    registered = {_team_identity(team): team for team in teams}
    source_set = {source for source, _target in mappings}
    for source, target in mappings:
        if source not in registered:
            raise SeasonStateError(
                "Refusing team rename: source identity is not in the registered season snapshot: "
                f"{source[0]}/{source[2]}/{source[1]}"
            )
        if target in registered and target not in source_set:
            raise SeasonStateError(
                "Refusing team rename: target identity already exists in the registered season snapshot: "
                f"{target[0]}/{target[2]}/{target[1]}"
            )


def _rename_team_refs(value: Any, mapping_by_source: Mapping[TeamIdentity, TeamIdentity]) -> tuple[Any, int]:
    """Recursively rename explicit ``club``/``label``/``age_group`` team references."""

    if isinstance(value, list):
        changed = 0
        items = []
        for item in value:
            new_item, item_changed = _rename_team_refs(item, mapping_by_source)
            changed += item_changed
            items.append(new_item)
        return items, changed
    if not isinstance(value, dict):
        return value, 0

    out: dict[str, Any] = {}
    changed = 0
    for key, item in value.items():
        new_item, item_changed = _rename_team_refs(item, mapping_by_source)
        changed += item_changed
        out[key] = new_item

    if {"club", "label", "age_group"}.issubset(out.keys()):
        identity = _team_identity(out)
        target = mapping_by_source.get(identity)
        if target is not None:
            out["club"], out["label"], out["age_group"] = target
            changed += 1
    return out, changed


def _rename_tournament_game_labels(
    plan: dict[str, Any],
    mapping_by_source: Mapping[TeamIdentity, TeamIdentity],
) -> int:
    changed = 0
    label_by_identity = {source: target[1] for source, target in mapping_by_source.items()}
    for tournament in plan.get("tournaments", []) or []:
        if not isinstance(tournament, dict):
            continue
        age_group = str(tournament.get("age_group") or "")
        labels = {}
        for source, target_label in label_by_identity.items():
            target_identity = mapping_by_source[source]
            if source[2] != age_group:
                continue
            if any(
                _team_identity(team, age_group) in {source, target_identity}
                for team in tournament.get("teams", []) or []
            ):
                labels[source[1]] = target_label
        if not labels:
            continue
        for game in tournament.get("games", []) or []:
            if not isinstance(game, dict):
                continue
            for field in ("home", "away"):
                current = str(game.get(field) or "")
                replacement = labels.get(current)
                if replacement is not None:
                    game[field] = replacement
                    changed += 1
    return changed


def _rename_context_problem(schedule: dict[str, Any], mapping_by_source: Mapping[TeamIdentity, TeamIdentity]) -> int:
    context = schedule.get("verification_context")
    if not isinstance(context, dict):
        return 0
    problem = context.get("problem")
    if not isinstance(problem, dict):
        return 0
    renamed_problem, changed = _rename_team_refs(problem, mapping_by_source)
    context["problem"] = renamed_problem
    context["problem_fingerprint"] = stable_payload_sha256(renamed_problem)
    promoted_from = schedule.get("promoted_from")
    if isinstance(promoted_from, dict):
        promoted_from["verification_context_problem_fingerprint"] = context["problem_fingerprint"]
    return changed


def _placement_signature(plan: Mapping[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        str(tournament.get("id") or ""): {
            "date": tournament.get("date"),
            "arena": tournament.get("arena"),
            "host_club": tournament.get("host_club"),
            "start_time": tournament.get("start_time"),
        }
        for tournament in plan.get("tournaments", []) or []
        if isinstance(tournament, Mapping) and tournament.get("id")
    }


def _participant_identity_signature(plan: Mapping[str, Any]) -> dict[str, list[TeamIdentity]]:
    signature: dict[str, list[TeamIdentity]] = {}
    for tournament in plan.get("tournaments", []) or []:
        if not isinstance(tournament, Mapping) or not tournament.get("id"):
            continue
        age_group = str(tournament.get("age_group") or "")
        signature[str(tournament.get("id"))] = [
            _team_identity(team, age_group)
            for team in tournament.get("teams", []) or []
            if isinstance(team, Mapping) and not bool(team.get("guest", False))
        ]
    return signature


def _rename_signature(
    signature: Mapping[str, Sequence[TeamIdentity]],
    mapping_by_source: Mapping[TeamIdentity, TeamIdentity],
) -> dict[str, list[TeamIdentity]]:
    return {
        tournament_id: [mapping_by_source.get(identity, identity) for identity in identities]
        for tournament_id, identities in signature.items()
    }


def _input_workbook_report(
    *,
    input_path: str | None,
    mappings: Sequence[RenameMapping],
    apply: bool,
) -> dict[str, Any]:
    report = {
        "path": input_path,
        "enabled": bool(input_path),
        "updated": False,
        "renamed_rows": [],
    }
    if not input_path:
        return report

    from openpyxl import load_workbook

    workbook = load_workbook(input_path)
    if "Lag" not in workbook.sheetnames:
        raise SeasonStateError(f"Input workbook {input_path!r} has no Lag sheet")
    sheet = workbook["Lag"]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    try:
        club_col = headers.index("club") + 1
        label_col = headers.index("label") + 1
        age_col = headers.index("age_group") + 1
    except ValueError as exc:
        raise SeasonStateError("Lag sheet must contain club, label and age_group columns") from exc
    mapping_by_source = dict(mappings)
    seen_sources: set[TeamIdentity] = set()
    for row_index in range(2, sheet.max_row + 1):
        identity = (
            str(sheet.cell(row_index, club_col).value or ""),
            str(sheet.cell(row_index, label_col).value or ""),
            str(sheet.cell(row_index, age_col).value or ""),
        )
        target = mapping_by_source.get(identity)
        if target is None:
            continue
        seen_sources.add(identity)
        report["renamed_rows"].append(
            {"row": row_index, "from": _team_ref(identity), "to": _team_ref(target)}
        )
        if apply:
            sheet.cell(row_index, label_col).value = target[1]
            report["updated"] = True
    missing = [source for source, _target in mappings if source not in seen_sources]
    if missing:
        formatted = ", ".join(f"{club}/{age}/{label}" for club, label, age in missing)
        raise SeasonStateError(f"Input workbook is missing source team identity/identities: {formatted}")
    if apply and report["updated"]:
        workbook.save(input_path)
    return report


def rename_teams(
    service,
    *,
    season: str,
    mappings: Sequence[Mapping[str, Any]],
    problem: dict[str, Any] | None = None,
    actor: str | None = None,
    note: str = "",
    dry_run: bool = False,
    request_id: str | None = None,
    input_path: str | None = None,
) -> dict[str, Any]:
    """Atomically migrate canonical team identities without changing placements."""

    normalized_mappings = _normalize_mappings(mappings)
    mapping_by_source = dict(normalized_mappings)
    snapshot = service.load(season)
    schedule, decisions = snapshot.schedule, snapshot.decisions
    resolved_problem = _resolve_plan_problem(schedule, problem, decisions)
    if resolved_problem is None:
        raise SeasonStateError("Team rename requires a promoted verification-context problem")
    _assert_registered_pool_allows_rename(resolved_problem, normalized_mappings)

    before_plan = schedule.get("plan") or {}
    before_ids = [str(t.get("id") or "") for t in before_plan.get("tournaments", []) or []]
    before_placements = _placement_signature(before_plan)
    before_participants = _participant_identity_signature(before_plan)
    before_revision = canonical_state_revision(schedule, decisions)

    plan, team_ref_changes = _rename_team_refs(copy.deepcopy(before_plan), mapping_by_source)
    if not isinstance(plan, dict):  # defensive; before_plan is dict in canonical state
        raise SeasonStateError("Canonical plan has invalid structure")
    game_changes = _rename_tournament_game_labels(plan, mapping_by_source)
    updated_schedule = copy.deepcopy(schedule)
    updated_decisions, decision_ref_changes = _rename_team_refs(copy.deepcopy(decisions), mapping_by_source)
    if not isinstance(updated_decisions, dict):
        raise SeasonStateError("Canonical decisions have invalid structure")
    context_ref_changes = _rename_context_problem(updated_schedule, mapping_by_source)

    # Re-resolve against the updated schedule context so projections and renamed registered teams match the plan.
    renamed_problem = _resolve_plan_problem(updated_schedule, None, updated_decisions)
    if renamed_problem is None:
        raise SeasonStateError("Team rename lost the promoted verification-context problem")
    result = verify_candidate(plan, renamed_problem)
    if not result.get("ok", True):
        messages = "; ".join(str(v.get("message") or v.get("code")) for v in result.get("violations", []))
        raise SeasonStateError(f"Refusing canonical team rename: candidate fails hard verification: {messages}")
    reconcile_plan_derived_state(plan, result, problem=renamed_problem)

    after_ids = [str(t.get("id") or "") for t in plan.get("tournaments", []) or []]
    after_placements = _placement_signature(plan)
    after_participants = _participant_identity_signature(plan)
    expected_participants = _rename_signature(before_participants, mapping_by_source)
    placement_neutral = {
        "tournaments_added": sorted(set(after_ids) - set(before_ids)),
        "tournaments_removed": sorted(set(before_ids) - set(after_ids)),
        "tournament_order_changed": before_ids != after_ids,
        "placement_changed_ids": sorted(
            tournament_id
            for tournament_id in set(before_placements) | set(after_placements)
            if before_placements.get(tournament_id) != after_placements.get(tournament_id)
        ),
        "participant_assignments_changed_except_rename": after_participants != expected_participants,
    }
    if any(placement_neutral.values()):
        raise SeasonStateError(
            "Refusing canonical team rename: mutation is not placement/assignment neutral: "
            + str(placement_neutral)
        )

    now = _now_iso()
    fingerprint = schedule_fingerprint(plan)
    plan["schema_version"] = SEASON_PLAN_SCHEMA_VERSION
    updated_schedule.update(
        {
            "updated_at": now,
            "revision": fingerprint,
            "fingerprint": fingerprint,
            "plan_schema_version": SEASON_PLAN_SCHEMA_VERSION,
            "plan": plan,
            "applied_from": {
                "previous_revision": schedule.get("revision"),
                "actor": _operator_identity(actor),
                "operation": "team_identity_rename",
            },
        }
    )
    approval_fingerprint_updates = 0
    tournaments_by_id = {
        str(tournament.get("id") or ""): tournament
        for tournament in plan.get("tournaments", []) or []
        if isinstance(tournament, Mapping) and tournament.get("id")
    }
    for tournament_id, record in (updated_decisions.get("decisions") or {}).items():
        if not isinstance(record, dict) or not record.get("approved_fingerprint"):
            continue
        tournament = tournaments_by_id.get(str(tournament_id))
        if tournament is None:
            continue
        record["approved_fingerprint"] = approval_fingerprint(tournament)
        approval_fingerprint_updates += 1

    updated_decisions.update(
        {
            "updated_at": now,
            "schedule_fingerprint": fingerprint,
        }
    )
    history = updated_decisions.setdefault("history", [])
    history.append(
        {
            "event": "team_identity_rename",
            "tournament_id": "",
            "actor": _operator_identity(actor),
            "at": now,
            "tournament_fingerprint": None,
            "previous_fingerprint": schedule.get("fingerprint"),
            "schedule_fingerprint": fingerprint,
            "note": note or "",
            "details": {
                "request_id": request_id or "",
                "mappings": [{"from": _team_ref(source), "to": _team_ref(target)} for source, target in normalized_mappings],
                "before_canonical_revision": before_revision,
                "renamed_team_identities": len(normalized_mappings),
                "team_reference_changes": team_ref_changes,
                "game_label_changes": game_changes,
                "decision_reference_changes": decision_ref_changes,
                "verification_context_reference_changes": context_ref_changes,
                "approval_fingerprint_updates": approval_fingerprint_updates,
                "placement_neutrality": placement_neutral,
            },
        }
    )

    workbook_report = _input_workbook_report(
        input_path=input_path,
        mappings=normalized_mappings,
        apply=not dry_run,
    )
    report = {
        "season": season,
        "dry_run": dry_run,
        "request_id": request_id or "",
        "renamed_team_identities": len(normalized_mappings),
        "mappings": [{"from": _team_ref(source), "to": _team_ref(target)} for source, target in normalized_mappings],
        "current_revision": schedule.get("revision"),
        "candidate_revision": fingerprint,
        "candidate_fingerprint": fingerprint,
        "before_canonical_revision": before_revision,
        "verification_result": result,
        "placement_neutrality": placement_neutral,
        "reference_changes": {
            "plan_team_refs": team_ref_changes,
            "game_labels": game_changes,
            "decisions": decision_ref_changes,
            "verification_context": context_ref_changes,
            "approval_fingerprints": approval_fingerprint_updates,
        },
        "input_workbook": workbook_report,
    }
    if dry_run:
        return report

    committed = service._commit(
        CanonicalSeasonSnapshot(
            season=snapshot.season,
            schedule=updated_schedule,
            decisions=updated_decisions,
            export_context=snapshot.export_context,
        )
    )
    report["revision"] = committed.schedule.get("revision")
    report["canonical_state_revision"] = canonical_state_revision(committed.schedule, committed.decisions)
    return report
