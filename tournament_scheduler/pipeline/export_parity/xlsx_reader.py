"""Read the on-disk ``season_plan.xlsx`` back into normalized records.

The reader is intentionally tolerant of the legacy workbook shape (no stable
id/status columns): it records those fields as *unsupported* so parity is
reported ``NOT_CHECKABLE`` instead of silently passing. It never trusts the
exporter's in-memory object.
"""

from __future__ import annotations

import hashlib
from dataclasses import replace
from pathlib import Path
from typing import Any

from .records import (
    ArtifactProjection,
    TournamentRecord,
    has_cancel_prefix,
    normalize_game,
    normalize_games,
    normalize_iso_date,
    normalize_participants,
    normalize_text,
    normalize_time,
    yes_no,
)

OVERVIEW_SHEET = "Sesongoversikt"
METADATA_SHEET = "Eksportmetadata"

_REQUIRED_COLUMNS = ("Dato", "Aldersgruppe", "Arena", "Vertsklubb", "Lag", "Starttid", "Sluttid")
_ID_COLUMN = "Turnerings-ID"
# XLSX column -> canonical parity field name. ``Avlyst`` is intentionally
# omitted: the cancellation flag is always derivable from the ``(AVLYST)`` date
# prefix, so it stays checkable even in a legacy workbook without the column.
_STATUS_COLUMNS = {
    "Avlysningsårsak": "cancellation_reason",
    "Godkjenning": "approval_status",
    "Låst": "locked",
    "Bookingsstatus": "booking_status",
    "Booking krever oppfølging": "booking_needs_attention",
}
_REVISION_LABEL = "Kanonisk revisjon"
_ID_ROW_PREFIX = "Turnerings-ID:"
_GAMES_HEADER = ("Runde", "Hjemmelag", "Bortelag", "Parallellbane")


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def _column_index(header: tuple[Any, ...]) -> dict[str, int]:
    index: dict[str, int] = {}
    for position, name in enumerate(header):
        key = normalize_text(name)
        if key and key not in index:
            index[key] = position
    return index


def _cell(row: tuple[Any, ...], columns: dict[str, int], name: str) -> Any:
    position = columns.get(name)
    if position is None or position >= len(row):
        return None
    return row[position]


def _read_metadata(workbook: Any) -> str:
    if METADATA_SHEET not in workbook.sheetnames:
        return ""
    sheet = workbook[METADATA_SHEET]
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        key = normalize_text(row[0])
        if key == _REVISION_LABEL:
            return normalize_text(row[1] if len(row) > 1 else "")
    return ""


def read_xlsx(path: str | Path) -> ArtifactProjection:
    file_path = Path(path)
    projection = ArtifactProjection(kind="xlsx", path=str(file_path), exists=file_path.exists())
    if not file_path.exists():
        return projection
    try:
        import openpyxl

        projection.sha256 = _file_sha256(file_path)
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            if OVERVIEW_SHEET not in workbook.sheetnames:
                projection.read_error = f"worksheet {OVERVIEW_SHEET!r} is missing"
                return projection
            rows = list(workbook[OVERVIEW_SHEET].iter_rows(values_only=True))
            if not rows:
                projection.read_error = "overview worksheet is empty"
                return projection
            columns = _column_index(tuple(rows[0]))
            missing_required = [name for name in _REQUIRED_COLUMNS if name not in columns]
            if missing_required:
                projection.read_error = (
                    "overview worksheet is missing required column(s): "
                    + ", ".join(missing_required)
                )
                return projection
            unsupported: list[str] = []
            if _ID_COLUMN not in columns:
                unsupported.append("tournament_id")
            for column, field in _STATUS_COLUMNS.items():
                if column not in columns:
                    unsupported.append(field)
            projection.unsupported_fields = unsupported
            projection.season_revision = _read_metadata(workbook)
            records = _read_rows(rows[1:], columns)
            games_by_id = _read_games_sheets(workbook, unsupported)
            projection.records = [
                replace(record, games=games_by_id.get(record.tournament_id, ())) for record in records
            ]
        finally:
            workbook.close()
    except Exception as exc:  # noqa: BLE001 - any unreadable bytes mean NOT_CHECKABLE
        projection.read_error = f"could not read workbook: {exc}"
    return projection


def _read_rows(rows: list[tuple[Any, ...]], columns: dict[str, int]) -> list[TournamentRecord]:
    records: list[TournamentRecord] = []
    for row in rows:
        if row is None or not any(cell not in (None, "") for cell in row):
            continue
        raw_date = _cell(row, columns, "Dato")
        cancelled = has_cancel_prefix(raw_date) or yes_no(_cell(row, columns, "Avlyst"))
        reason = normalize_text(_cell(row, columns, "Avlysningsårsak"))
        records.append(
            TournamentRecord(
                tournament_id=normalize_text(_cell(row, columns, _ID_COLUMN)),
                date=normalize_iso_date(raw_date),
                start_time=normalize_time(_cell(row, columns, "Starttid")),
                end_time=normalize_time(_cell(row, columns, "Sluttid")),
                arena=normalize_text(_cell(row, columns, "Arena")),
                host_club=normalize_text(_cell(row, columns, "Vertsklubb")),
                age_group=normalize_text(_cell(row, columns, "Aldersgruppe")),
                participants=normalize_participants(
                    part for part in normalize_text(_cell(row, columns, "Lag")).split(",")
                ),
                cancelled=cancelled,
                cancellation_reason=reason,
                approval_status=normalize_text(_cell(row, columns, "Godkjenning")),
                locked=yes_no(_cell(row, columns, "Låst")),
                booking_status=normalize_text(_cell(row, columns, "Bookingsstatus")),
                booking_needs_attention=yes_no(_cell(row, columns, "Booking krever oppfølging")),
            )
        )
    return records


def _read_games_sheets(workbook: Any, unsupported: list[str]) -> dict[str, tuple[str, ...]]:
    """Read each per-tournament game sheet by its stable ``Turnerings-ID`` row.

    The sheet *title* is not a reliable identity (two same-day/same-host/same-age
    tournaments collide), so the id is read from the sheet body. A new-format
    workbook always carries it; if the id row is absent the format genuinely
    cannot expose game parity and ``games`` is reported uncheckable.
    """

    games_by_id: dict[str, tuple[str, ...]] = {}
    found_any_id = False
    for name in workbook.sheetnames:
        if name in {OVERVIEW_SHEET, METADATA_SHEET}:
            continue
        rows = list(workbook[name].iter_rows(values_only=True))
        tournament_id = ""
        for row in rows:
            if row and normalize_text(row[0]).startswith(_ID_ROW_PREFIX):
                tournament_id = normalize_text(row[0])[len(_ID_ROW_PREFIX):].strip()
                break
        header_index = next(
            (index for index, row in enumerate(rows) if tuple(row[:4]) == _GAMES_HEADER),
            None,
        )
        if header_index is None:
            continue
        found_any_id = found_any_id or bool(tournament_id)
        games: list[str] = []
        for row in rows[header_index + 1:]:
            if not row or not isinstance(row[0], int):
                continue
            if normalize_text(row[1]) == "(Pause)":
                continue
            try:
                slot = int(row[3]) - 1  # exporter displays 1-based slot
            except (TypeError, ValueError):
                slot = 0
            games.append(normalize_game(row[0], row[1], row[2], slot))
        if tournament_id:
            games_by_id[tournament_id] = normalize_games(games)
    if not found_any_id and "tournament_id" not in unsupported:
        unsupported.append("games")
    return games_by_id
