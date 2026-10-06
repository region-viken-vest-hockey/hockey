"""Independent reader for the written Spond season-plan import workbook.

The authoritative Spond projection is the ``Spond import`` sheet, not the
one-row summary sheet. Every non-empty data row is one participant in one
active tournament. Rows are grouped by stable ``RVV-ID`` and all placement
values must agree within a group. Duplicate participant rows and rows without a
stable identity remain structured diagnostics so malformed import bytes cannot
collapse into an apparently valid set projection.
"""

from __future__ import annotations

import hashlib
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from tournament_scheduler.spond.spond_import import SPOND_IMPORT_SHEET

from .records import (
    ArtifactProjection,
    TournamentRecord,
    normalize_iso_date,
    normalize_participant_keys,
    normalize_participants,
    normalize_text,
    normalize_time,
    participant_identity,
)

METADATA_SHEET = "Eksportmetadata"
_REVISION_LABEL = "Kanonisk revisjon"
_REQUIRED_COLUMNS = (
    "Startdato*",
    "Starttidspunkt",
    "Sluttidspunkt",
    "Sted",
    "Lag",
    "Klubb",
    "Aldersgruppe",
    "Vertsklubb",
    "RVV-ID",
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _text_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return normalize_iso_date(value)


def _text_time(value: Any) -> str:
    if isinstance(value, datetime):
        return value.time().strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return normalize_time(value)


def _columns(header: tuple[Any, ...]) -> dict[str, int]:
    return {
        normalize_text(value): index
        for index, value in enumerate(header)
        if normalize_text(value)
    }


def _cell(row: tuple[Any, ...], columns: dict[str, int], name: str) -> Any:
    index = columns.get(name)
    return row[index] if index is not None and index < len(row) else None


def _revision(workbook: Any) -> str:
    if METADATA_SHEET not in workbook.sheetnames:
        return ""
    for row in workbook[METADATA_SHEET].iter_rows(values_only=True):
        if row and normalize_text(row[0]) == _REVISION_LABEL:
            return normalize_text(row[1] if len(row) > 1 else "")
    return ""


def read_spond(path: str | Path) -> ArtifactProjection:
    file_path = Path(path)
    result = ArtifactProjection(kind="spond", path=str(file_path), exists=file_path.exists())
    if not file_path.exists():
        return result
    try:
        import openpyxl

        result.sha256 = _sha256(file_path)
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            if SPOND_IMPORT_SHEET not in workbook.sheetnames:
                result.read_error = f"worksheet {SPOND_IMPORT_SHEET!r} is missing"
                return result
            rows = list(workbook[SPOND_IMPORT_SHEET].iter_rows(values_only=True))
            if not rows:
                result.read_error = "Spond import worksheet is empty"
                return result
            columns = _columns(tuple(rows[0]))
            missing = [name for name in _REQUIRED_COLUMNS if name not in columns]
            if missing:
                result.read_error = "Spond import worksheet is missing required column(s): " + ", ".join(missing)
                return result
            result.season_revision = _revision(workbook)
            result.records, result.issues, result.row_count = _read_rows(rows[1:], columns)
        finally:
            workbook.close()
    except Exception as exc:  # noqa: BLE001 - malformed bytes must fail closed
        result.read_error = f"could not read workbook: {exc}"
    return result


def _read_rows(
    rows: list[tuple[Any, ...]], columns: dict[str, int]
) -> tuple[list[TournamentRecord], list[dict[str, Any]], int]:
    grouped: dict[str, list[dict[str, str]]] = {}
    issues: list[dict[str, Any]] = []
    row_count = 0
    for excel_row, row in enumerate(rows, start=2):
        if not row or not any(value not in (None, "") for value in row):
            continue
        row_count += 1
        tournament_id = normalize_text(_cell(row, columns, "RVV-ID"))
        participant = participant_identity(
            _cell(row, columns, "Klubb"),
            _cell(row, columns, "Lag"),
            _cell(row, columns, "Aldersgruppe"),
        )
        if not tournament_id:
            issues.append({
                "code": "unmatched_spond_row",
                "row": excel_row,
                "participant": participant,
                "message": "Spond row has no RVV-ID",
            })
            continue
        grouped.setdefault(tournament_id, []).append({
            "date": _text_date(_cell(row, columns, "Startdato*")),
            "start_time": _text_time(_cell(row, columns, "Starttidspunkt")),
            "end_time": _text_time(_cell(row, columns, "Sluttidspunkt")),
            "arena": normalize_text(_cell(row, columns, "Sted")),
            "host_club": normalize_text(_cell(row, columns, "Vertsklubb")),
            "age_group": normalize_text(_cell(row, columns, "Aldersgruppe")),
            "participant": participant,
            "participant_label": normalize_text(_cell(row, columns, "Lag")),
            "row": str(excel_row),
        })

    records: list[TournamentRecord] = []
    placement_fields = ("date", "start_time", "end_time", "arena", "host_club", "age_group")
    for tournament_id, group in sorted(grouped.items()):
        first = group[0]
        counts = Counter(item["participant"] for item in group)
        for participant, count in sorted(counts.items()):
            if count > 1:
                issues.append({
                    "code": "duplicate_spond_row",
                    "tournament_id": tournament_id,
                    "participant": participant,
                    "count": count,
                    "rows": [int(item["row"]) for item in group if item["participant"] == participant],
                    "message": "Spond contains more than one import row for this participant",
                })
        for field in placement_fields:
            values = sorted({item[field] for item in group})
            if len(values) > 1:
                issues.append({
                    "code": "inconsistent_spond_tournament",
                    "tournament_id": tournament_id,
                    "field": field,
                    "values": values,
                    "message": "Spond rows for one tournament disagree",
                })
        records.append(TournamentRecord(
            tournament_id=tournament_id,
            date=first["date"],
            start_time=first["start_time"],
            end_time=first["end_time"],
            arena=first["arena"],
            host_club=first["host_club"],
            age_group=first["age_group"],
            participants=normalize_participants(item["participant_label"] for item in group),
            participant_keys=normalize_participant_keys(item["participant"] for item in group),
        ))
    return records, issues, row_count


__all__ = ["read_spond"]
