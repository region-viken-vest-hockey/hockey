"""Schedule-preservation guard for canonical season exports.

A promoted season export is a projection of canonical state.  This module keeps
that invariant independent from renderers and export formats by comparing the
stable tournament-id projection before Stage 4 writes artifacts.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
import json
from pathlib import Path
import re
from typing import Any, Mapping

from tournament_scheduler.guest_slots import (
    guest_slot_records,
)


@dataclass(frozen=True)
class ExportProjectionError(ValueError):
    """Raised when an export would change canonical schedule semantics."""

    report: dict[str, Any]

    def __str__(self) -> str:
        summary = self.report.get("summary") or "season export would change canonical schedule"
        return str(summary)


FULL_OPERATIONAL_PROJECTION_SCHEMA = "rvv.operational_tournament_projection"
FULL_OPERATIONAL_PROJECTION_VERSION = 2
LEGACY_ARTIFACT_BACKFILL = "legacy_artifact_backfill_from_publication_revision"

_OPERATIONAL_FIELDS: tuple[str, ...] = (
    "id",
    "age_group",
    "date",
    "start_time",
    "arena",
    "host_club",
    "duration_minutes",
    "end_time",
    "cancelled",
    "cancellation_reason",
    "participants",
    "guest_slots",
)
_REQUIRED_ENTRY_FIELDS = {
    "projection_schema",
    "projection_schema_version",
    *_OPERATIONAL_FIELDS,
}


def _participant_key(team: Mapping[str, Any]) -> str:
    return "\u001f".join(
        str(team.get(field) or "") for field in ("club", "label", "age_group")
    )


def _team_identity(team: Any, *, fallback_age_group: str = "") -> dict[str, str]:
    if not isinstance(team, Mapping):
        return {"club": "", "label": "", "age_group": fallback_age_group}
    return {
        "club": str(team.get("club") or ""),
        "label": str(team.get("label") or ""),
        "age_group": str(team.get("age_group") or fallback_age_group),
    }


def _normalized_guest_slots(tournament: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Return only operational guest-reservation facts.

    Reservation actor/timestamp/note fields are audit evidence. They are kept in
    canonical history, but they are not schedule facts and therefore are not
    part of the publication-preservation projection.
    """

    age_group = str(tournament.get("age_group") or "")
    slots: list[dict[str, Any]] = []
    for record in guest_slot_records(tournament):
        if not isinstance(record, Mapping):
            continue
        slots.append(
            {
                "id": str(record.get("id") or ""),
                "status": str(record.get("status") or "open"),
                "external_team": _team_identity(
                    record.get("external_team"), fallback_age_group=age_group
                )
                if isinstance(record.get("external_team"), Mapping)
                else None,
            }
        )
    return sorted(slots, key=lambda item: (item["id"], item["status"], str(item["external_team"])))


def _parse_hhmm(value: Any) -> datetime | None:
    try:
        return datetime.strptime(str(value), "%H:%M")
    except (TypeError, ValueError):
        return None


def occupied_end_time(start_time: Any, duration_minutes: Any) -> str:
    """Canonical end time for a start time plus occupied duration."""

    try:
        minutes = int(duration_minutes)
    except (TypeError, ValueError):
        return ""
    parsed = _parse_hhmm(start_time)
    if parsed is None or minutes <= 0:
        return ""
    return (parsed + timedelta(minutes=minutes)).strftime("%H:%M")


def _end_time(start_time: Any, duration_minutes: int) -> str:
    return occupied_end_time(start_time, duration_minutes)


def _duration_from_problem(tournament: Mapping[str, Any], problem: Mapping[str, Any] | None) -> int | None:
    age_group = str(tournament.get("age_group") or "")
    ice_time = (problem or {}).get("ice_time_minutes") or {}
    if isinstance(ice_time, Mapping) and age_group in ice_time:
        try:
            duration = int(ice_time[age_group])
        except (TypeError, ValueError):
            return None
        return duration if duration > 0 else None
    for field in ("duration_minutes", "ice_time_minutes"):
        if field in tournament:
            try:
                duration = int(tournament[field])
            except (TypeError, ValueError):
                return None
            return duration if duration > 0 else None
    return None


def _projection_entry_from_tournament(
    tournament: Mapping[str, Any],
    *,
    problem: Mapping[str, Any] | None,
    artifact_overlay: Mapping[str, Any] | None = None,
    migration: str | None = None,
    strict: bool = True,
) -> dict[str, Any]:
    tournament_id = str((artifact_overlay or {}).get("id") or tournament.get("id") or "")
    age_group = str((artifact_overlay or {}).get("age_group") or tournament.get("age_group") or "")
    duration = _duration_from_problem({**dict(tournament), "age_group": age_group}, problem)
    if duration is None:
        if strict:
            raise ExportProjectionError(
                {
                    "summary": (
                        "Refusing schedule projection: canonical occupied duration is unknown for "
                        f"tournament {tournament_id or '<missing id>'} ({age_group or '<missing age group>'})"
                    ),
                    "tournament_id": tournament_id,
                    "age_group": age_group,
                    "required_schema": FULL_OPERATIONAL_PROJECTION_SCHEMA,
                }
            )
        duration = 0
    start_time = str((artifact_overlay or {}).get("start_time") or tournament.get("start_time") or "")
    teams = (artifact_overlay or {}).get("teams") or tournament.get("teams") or []
    entry: dict[str, Any] = {
        "projection_schema": FULL_OPERATIONAL_PROJECTION_SCHEMA,
        "projection_schema_version": FULL_OPERATIONAL_PROJECTION_VERSION,
        "id": tournament_id,
        "date": str((artifact_overlay or {}).get("date") or tournament.get("date") or ""),
        "start_time": start_time,
        "arena": str((artifact_overlay or {}).get("arena") or tournament.get("arena") or ""),
        "host_club": str((artifact_overlay or {}).get("host_club") or tournament.get("host_club") or ""),
        "age_group": age_group,
        "duration_minutes": duration,
        "end_time": _end_time(start_time, duration),
        "cancelled": bool(tournament.get("cancelled", False)),
        "cancellation_reason": str(tournament.get("cancellation_reason") or ""),
        "participants": sorted(
            _participant_key(team)
            for team in teams
            if isinstance(team, Mapping)
        ),
        "guest_slots": _normalized_guest_slots(tournament),
    }
    if migration:
        entry["projection_migration"] = migration
    return entry


def tournament_projection(
    plan: Mapping[str, Any] | None,
    problem: Mapping[str, Any] | None = None,
    *,
    strict: bool = True,
) -> dict[str, dict[str, Any]]:
    """Return the versioned full operational schedule projection.

    The projection intentionally contains schedule facts only: stable identity,
    placement, roster identity, canonical occupied interval, cancellation state
    and guest-reservation facts. Approval, booking evidence, audit provenance and
    other evidence-only decisions stay outside this predicate.

    ``strict`` is the preservation-check mode: an unknown occupied duration
    fails closed because the projection cannot be trusted. A non-strict
    projection is only used for a non-canonical diagnostic manifest, where the
    duration is left as zero rather than invented.
    """

    projection: dict[str, dict[str, Any]] = {}
    for tournament in (plan or {}).get("tournaments", []) or []:
        if not isinstance(tournament, Mapping):
            continue
        tournament_id = str(tournament.get("id") or "")
        if not tournament_id:
            continue
        projection[tournament_id] = _projection_entry_from_tournament(
            tournament,
            problem=problem,
            strict=strict,
        )
    return projection


_TOURNAMENTS_RE = re.compile(r"^\s*const\s+TOURNAMENTS\s*=\s*(\[.*\]);\s*$", re.MULTILINE)


def projection_from_export_artifacts(
    export_dir: str | Path,
    *,
    canonical_plan: Mapping[str, Any] | None = None,
    published_canonical_plan: Mapping[str, Any] | None = None,
    published_canonical_problem: Mapping[str, Any] | None = None,
) -> dict[str, dict[str, Any]]:
    """Reconstruct a stable-id projection from legacy published artifacts.

    Manifests written before the publication guard did not carry
    ``schedule_projection``.  Prefer the embedded HTML tournament payload when
    present; it carries stable ids directly.  The committed 2026-09-21
    publication did not track that HTML, but it does track the Spond season
    workbook, which carries placement/participants but no stable id.  For that
    legacy shape the row-to-id identity must be recovered from the canonical
    plan *at the publication revision* (``published_canonical_plan``); binding
    against the current canonical plan or against row order would fabricate
    identity.  Refuse partial reconstruction: publication safety must not
    silently compare against an incomplete or guessed baseline.
    """

    html_path = Path(export_dir) / "season_plan.html"
    try:
        text = html_path.read_text(encoding="utf-8")
    except OSError:
        return _projection_from_spond_workbook(
            export_dir,
            canonical_plan=canonical_plan,
            published_canonical_plan=published_canonical_plan,
            published_canonical_problem=published_canonical_problem,
        )
    match = _TOURNAMENTS_RE.search(text)
    if not match:
        raise ExportProjectionError(
            {
                "summary": (
                    "Refusing season export: published export has no schedule_projection "
                    f"and {html_path} does not contain embedded TOURNAMENTS data"
                ),
                "export_dir": str(export_dir),
            }
        )
    try:
        rendered = json.loads(match.group(1))
    except json.JSONDecodeError as exc:
        raise ExportProjectionError(
            {
                "summary": (
                    "Refusing season export: published export has no schedule_projection "
                    f"and {html_path} contains invalid embedded TOURNAMENTS JSON"
                ),
                "export_dir": str(export_dir),
            }
        ) from exc
    if not isinstance(rendered, list):
        raise ExportProjectionError(
            {
                "summary": "Refusing season export: legacy TOURNAMENTS payload is not a list",
                "export_dir": str(export_dir),
            }
        )
    canonical_by_id = {
        str(tournament.get("id") or ""): tournament
        for tournament in (published_canonical_plan or {}).get("tournaments", []) or []
        if isinstance(tournament, Mapping) and tournament.get("id")
    }
    projection: dict[str, dict[str, Any]] = {}
    for item in rendered:
        if not isinstance(item, Mapping):
            continue
        tournament_id = str(item.get("id") or "")
        if not tournament_id:
            continue
        canonical_tournament = canonical_by_id.get(tournament_id)
        if canonical_tournament is None:
            raise ExportProjectionError(
                {
                    "summary": (
                        "Refusing season export: legacy TOURNAMENTS payload needs "
                        "the publication-time canonical schedule to backfill the "
                        "versioned operational projection"
                    ),
                    "tournament_id": tournament_id,
                    "export_dir": str(export_dir),
                }
            )
        participants = item.get("p") or []
        if not isinstance(participants, list):
            participants = []
        artifact_overlay = {
            "id": tournament_id,
            "date": str(item.get("d") or ""),
            "start_time": str(item.get("ts") or ""),
            "arena": str(item.get("a") or ""),
            "host_club": str(item.get("h") or ""),
            "age_group": str(item.get("g") or ""),
            "teams": [
                {"club": team.get("c"), "label": team.get("l"), "age_group": team.get("g")}
                for team in participants
                if isinstance(team, Mapping)
            ],
        }
        projection[tournament_id] = _projection_entry_from_tournament(
            canonical_tournament,
            problem=published_canonical_problem,
            artifact_overlay=artifact_overlay,
            migration=LEGACY_ARTIFACT_BACKFILL,
        )
    if not projection:
        raise ExportProjectionError(
            {
                "summary": "Refusing season export: legacy projection reconstruction produced no tournaments",
                "export_dir": str(export_dir),
            }
        )
    return projection


def _projection_from_spond_workbook(
    export_dir: str | Path,
    *,
    canonical_plan: Mapping[str, Any] | None,
    published_canonical_plan: Mapping[str, Any] | None,
    published_canonical_problem: Mapping[str, Any] | None,
) -> dict[str, dict[str, Any]]:
    workbook_path = Path(export_dir) / "season_plan_spond.xlsx"
    if not published_canonical_plan:
        raise ExportProjectionError(
            {
                "summary": (
                    "Refusing season export: published export has no schedule_projection, "
                    "HTML fallback is unavailable, and the canonical schedule snapshot "
                    "from the publication revision could not be resolved; refusing to bind "
                    "legacy Spond workbook rows to stable tournament ids by row order or "
                    "against post-publication canonical state"
                ),
                "export_dir": str(export_dir),
            }
        )
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is installed in normal runtime
        raise ExportProjectionError(
            {
                "summary": (
                    "Refusing season export: published export has no schedule_projection "
                    "and openpyxl is unavailable for legacy Spond workbook reconstruction"
                ),
                "export_dir": str(export_dir),
            }
        ) from exc
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except OSError as exc:
        raise ExportProjectionError(
            {
                "summary": (
                    "Refusing season export: published export has no schedule_projection "
                    f"and legacy projection artifacts are unreadable: {workbook_path}"
                ),
                "export_dir": str(export_dir),
            }
        ) from exc
    if "Sesongplan" not in workbook.sheetnames:
        raise ExportProjectionError(
            {
                "summary": "Refusing season export: legacy Spond workbook has no Sesongplan sheet",
                "export_dir": str(export_dir),
            }
        )
    rows = list(workbook["Sesongplan"].iter_rows(values_only=True))
    if not rows:
        raise ExportProjectionError(
            {
                "summary": "Refusing season export: legacy Spond workbook is empty",
                "export_dir": str(export_dir),
            }
        )
    header = [str(value or "") for value in rows[0]]
    index = {name: idx for idx, name in enumerate(header)}
    required = ["Dato", "Sted", "Start", "Aldersgruppe", "Vertsklubb", "Deltakende lag"]
    missing = [name for name in required if name not in index]
    if missing:
        raise ExportProjectionError(
            {
                "summary": "Refusing season export: legacy Spond workbook lacks required columns",
                "missing_columns": missing,
                "export_dir": str(export_dir),
            }
        )

    published_tournaments = [
        tournament
        for tournament in (published_canonical_plan or {}).get("tournaments", []) or []
        if isinstance(tournament, Mapping) and tournament.get("id")
    ]
    if not published_tournaments:
        raise ExportProjectionError(
            {
                "summary": (
                    "Refusing season export: the canonical schedule snapshot from the "
                    "publication revision contains no stable tournament ids"
                ),
                "export_dir": str(export_dir),
            }
        )
    data_rows = [row for row in rows[1:] if any(cell not in (None, "") for cell in row)]

    def cell(row: tuple[Any, ...], name: str) -> str:
        value = row[index[name]] if index[name] < len(row) else ""
        return str(value or "").strip()

    def iso_date(value: str) -> str:
        parts = value.split(".")
        if len(parts) == 3:
            return f"{parts[2]}-{parts[1]}-{parts[0]}"
        return value

    def row_identity(row: tuple[Any, ...]) -> tuple[Any, ...]:
        participant_labels = tuple(
            sorted(
                label.strip()
                for label in cell(row, "Deltakende lag").split(",")
                if label.strip()
            )
        )
        return (
            iso_date(cell(row, "Dato")),
            cell(row, "Start"),
            cell(row, "Sted"),
            cell(row, "Vertsklubb"),
            cell(row, "Aldersgruppe"),
            participant_labels,
        )

    def tournament_identity(tournament: Mapping[str, Any]) -> tuple[Any, ...]:
        labels = tuple(
            sorted(
                str(team.get("label") or "").strip()
                for team in tournament.get("teams") or []
                if isinstance(team, Mapping) and str(team.get("label") or "").strip()
            )
        )
        return (
            str(tournament.get("date") or ""),
            str(tournament.get("start_time") or ""),
            str(tournament.get("arena") or ""),
            str(tournament.get("host_club") or ""),
            str(tournament.get("age_group") or ""),
            labels,
        )

    by_identity: dict[tuple[Any, ...], list[Mapping[str, Any]]] = {}
    for tournament in published_tournaments:
        by_identity.setdefault(tournament_identity(tournament), []).append(tournament)

    projection: dict[str, dict[str, Any]] = {}
    unbound_rows: list[dict[str, Any]] = []
    for row in data_rows:
        identity = row_identity(row)
        candidates = by_identity.get(identity) or []
        if len(candidates) != 1:
            unbound_rows.append(
                {
                    "date": identity[0],
                    "start_time": identity[1],
                    "arena": identity[2],
                    "age_group": identity[4],
                    "participant_labels": list(identity[5]),
                    "candidate_count": len(candidates),
                }
            )
            continue
        tournament = candidates[0]
        tournament_id = str(tournament.get("id") or "")
        if tournament_id in projection:
            unbound_rows.append(
                {
                    "date": identity[0],
                    "start_time": identity[1],
                    "arena": identity[2],
                    "age_group": identity[4],
                    "participant_labels": list(identity[5]),
                    "candidate_count": 0,
                    "reason": "stable_id_already_bound",
                    "tournament_id": tournament_id,
                }
            )
            continue
        artifact_overlay = {
            "id": tournament_id,
            "date": identity[0],
            "start_time": identity[1],
            "arena": identity[2],
            "host_club": identity[3],
            "age_group": identity[4],
            "teams": [
                team
                for team in tournament.get("teams") or []
                if isinstance(team, Mapping)
            ],
        }
        projection[tournament_id] = _projection_entry_from_tournament(
            tournament,
            problem=published_canonical_problem,
            artifact_overlay=artifact_overlay,
            migration=LEGACY_ARTIFACT_BACKFILL,
        )
    if unbound_rows:
        raise ExportProjectionError(
            {
                "summary": (
                    "Refusing season export: legacy Spond workbook rows cannot be "
                    "uniquely bound to stable tournament ids from the canonical "
                    "schedule snapshot at the publication revision"
                ),
                "unbound_rows": unbound_rows,
                "workbook_tournament_count": len(data_rows),
                "export_dir": str(export_dir),
            }
        )
    if not projection:
        raise ExportProjectionError(
            {
                "summary": "Refusing season export: legacy Spond reconstruction produced no tournaments",
                "export_dir": str(export_dir),
            }
        )
    return projection


def _published_schedule_projection(
    published_export: Mapping[str, Any] | None,
    *,
    canonical_plan: Mapping[str, Any] | None = None,
    published_canonical_plan: Mapping[str, Any] | None = None,
    published_canonical_problem: Mapping[str, Any] | None = None,
) -> Mapping[str, Any] | None:
    if not published_export:
        return None
    published_projection = published_export.get("schedule_projection")
    if isinstance(published_projection, Mapping):
        return published_projection
    export_dir = published_export.get("export_dir")
    if export_dir:
        return projection_from_export_artifacts(
            export_dir,
            canonical_plan=canonical_plan,
            published_canonical_plan=published_canonical_plan,
            published_canonical_problem=published_canonical_problem,
        )
    raise ExportProjectionError(
        {
            "summary": (
                "Refusing season export: published export has no schedule_projection "
                "and no export_dir for legacy reconstruction"
            ),
            "published_export": {
                key: published_export.get(key)
                for key in ("export_id", "export_fingerprint", "canonical_revision", "lifecycle_status")
                if published_export.get(key) is not None
            },
        }
    )


def _projection_schema_errors(
    projection: Mapping[str, Mapping[str, Any]],
    *,
    side: str,
) -> list[dict[str, Any]]:
    errors: list[dict[str, Any]] = []
    if not isinstance(projection, Mapping):
        return [{"side": side, "reason": "projection_not_mapping"}]
    for tournament_id, entry in projection.items():
        stable_id = str(tournament_id or "")
        if not stable_id:
            errors.append({"side": side, "reason": "empty_tournament_id"})
            continue
        if not isinstance(entry, Mapping):
            errors.append({"side": side, "tournament_id": stable_id, "reason": "entry_not_mapping"})
            continue
        schema = entry.get("projection_schema")
        version = entry.get("projection_schema_version")
        if schema != FULL_OPERATIONAL_PROJECTION_SCHEMA or version != FULL_OPERATIONAL_PROJECTION_VERSION:
            errors.append(
                {
                    "side": side,
                    "tournament_id": stable_id,
                    "reason": "unsupported_projection_schema",
                    "schema": schema,
                    "schema_version": version,
                    "required_schema": FULL_OPERATIONAL_PROJECTION_SCHEMA,
                    "required_schema_version": FULL_OPERATIONAL_PROJECTION_VERSION,
                }
            )
            continue
        missing = sorted(field for field in _REQUIRED_ENTRY_FIELDS if field not in entry)
        if missing:
            errors.append(
                {"side": side, "tournament_id": stable_id, "reason": "missing_fields", "fields": missing}
            )
        if str(entry.get("id") or "") != stable_id:
            errors.append(
                {
                    "side": side,
                    "tournament_id": stable_id,
                    "reason": "mismatched_entry_id",
                    "entry_id": entry.get("id"),
                }
            )
        if not isinstance(entry.get("participants"), list):
            errors.append(
                {"side": side, "tournament_id": stable_id, "reason": "participants_not_list"}
            )
        if not isinstance(entry.get("guest_slots"), list):
            errors.append({"side": side, "tournament_id": stable_id, "reason": "guest_slots_not_list"})
        if not isinstance(entry.get("cancelled"), bool):
            errors.append({"side": side, "tournament_id": stable_id, "reason": "cancelled_not_bool"})
        if not isinstance(entry.get("duration_minutes"), int) or int(entry.get("duration_minutes") or 0) <= 0:
            errors.append(
                {"side": side, "tournament_id": stable_id, "reason": "duration_minutes_not_positive_int"}
            )
    return errors


def _normalized_compare_value(value: Any) -> Any:
    if isinstance(value, list):
        normalized = [_normalized_compare_value(item) for item in value]
        return sorted(normalized, key=lambda item: json.dumps(item, ensure_ascii=False, sort_keys=True))
    if isinstance(value, Mapping):
        return {str(key): _normalized_compare_value(value[key]) for key in sorted(value)}
    return value


def diff_tournament_projection(
    before: Mapping[str, Mapping[str, Any]],
    after: Mapping[str, Mapping[str, Any]],
) -> dict[str, Any]:
    before_ids = set(before)
    after_ids = set(after)
    removed = sorted(before_ids - after_ids)
    added = sorted(after_ids - before_ids)
    schema_errors = [
        *_projection_schema_errors(before, side="before"),
        *_projection_schema_errors(after, side="after"),
    ]
    field_changes: list[dict[str, Any]] = []
    placement_changes: list[dict[str, Any]] = []
    participant_changes: list[dict[str, Any]] = []
    guest_slot_changes: list[dict[str, Any]] = []
    cancellation_changes: list[dict[str, Any]] = []
    duration_changes: list[dict[str, Any]] = []
    placement_fields = ("date", "start_time", "arena", "host_club")
    compare_fields = tuple(
        field
        for field in _OPERATIONAL_FIELDS
        if field != "id"
    )
    for tournament_id in sorted(before_ids & after_ids):
        old = before[tournament_id]
        new = after[tournament_id]
        changes = {
            field: {"before": old.get(field), "after": new.get(field)}
            for field in compare_fields
            if _normalized_compare_value(old.get(field)) != _normalized_compare_value(new.get(field))
        }
        if not changes:
            continue
        field_changes.append({"tournament_id": tournament_id, "fields": changes})
        placement = {field: changes[field] for field in placement_fields if field in changes}
        if placement:
            placement_changes.append({"tournament_id": tournament_id, "fields": placement})
        if "participants" in changes:
            participant_changes.append(
                {
                    "tournament_id": tournament_id,
                    "before": list(old.get("participants") or []),
                    "after": list(new.get("participants") or []),
                }
            )
        guest = {field: changes[field] for field in ("guest_slots",) if field in changes}
        if guest:
            guest_slot_changes.append({"tournament_id": tournament_id, "fields": guest})
        cancellation = {
            field: changes[field]
            for field in ("cancelled", "cancellation_reason")
            if field in changes
        }
        if cancellation:
            cancellation_changes.append({"tournament_id": tournament_id, "fields": cancellation})
        duration = {
            field: changes[field]
            for field in ("duration_minutes", "end_time")
            if field in changes
        }
        if duration:
            duration_changes.append({"tournament_id": tournament_id, "fields": duration})
    return {
        "removed_tournament_ids": removed,
        "added_tournament_ids": added,
        "schema_errors": schema_errors,
        "field_changes": field_changes,
        "placement_changes": placement_changes,
        "participant_changes": participant_changes,
        "guest_slot_changes": guest_slot_changes,
        "cancellation_changes": cancellation_changes,
        "duration_changes": duration_changes,
        "changed": bool(removed or added or schema_errors or field_changes),
    }


def assert_export_preserves_canonical_plan(
    *,
    canonical_plan: Mapping[str, Any],
    proposed_plan: Mapping[str, Any],
    canonical_problem: Mapping[str, Any] | None = None,
    proposed_problem: Mapping[str, Any] | None = None,
    season: str | None = None,
    canonical_revision: str | None = None,
    published_export: Mapping[str, Any] | None = None,
    published_canonical_plan: Mapping[str, Any] | None = None,
    published_canonical_problem: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Raise if a proposed export is not an exact projection of canonical state.

    ``published_canonical_plan`` is the canonical plan at the revision the
    published export recorded. Legacy published artifacts that carry no stable
    tournament identity can only be reconstructed against that publication-time
    snapshot; passing the current canonical plan instead would fabricate row
    identity from row order.
    """

    canonical = tournament_projection(canonical_plan, canonical_problem)
    proposed = tournament_projection(proposed_plan, proposed_problem or canonical_problem)
    canonical_delta = diff_tournament_projection(canonical, proposed)
    published_projection = _published_schedule_projection(
        published_export,
        canonical_plan=canonical_plan,
        published_canonical_plan=published_canonical_plan,
        published_canonical_problem=published_canonical_problem,
    )
    published_delta: dict[str, Any] | None = None
    if isinstance(published_projection, Mapping):
        published_delta = diff_tournament_projection(published_projection, canonical)
    report: dict[str, Any] = {
        "season": season,
        "canonical_revision": canonical_revision,
        "canonical_tournament_count": len(canonical),
        "proposed_tournament_count": len(proposed),
        "canonical_delta": canonical_delta,
        "published_export": {
            key: (published_export or {}).get(key)
            for key in ("export_id", "export_fingerprint", "canonical_revision", "lifecycle_status")
            if (published_export or {}).get(key) is not None
        },
        "published_to_canonical_delta": published_delta,
    }
    if canonical_delta["changed"]:
        report["summary"] = (
            "Refusing season export: proposed artifacts would remove, move, "
            "or change operational tournament facts relative to the canonical schedule. "
            "Resolve conflicts through an explicit canonical season operation; "
            "export never repairs a published schedule."
        )
        raise ExportProjectionError(report)
    report["summary"] = "export projection preserves canonical schedule"
    return report


__all__ = [
    "ExportProjectionError",
    "assert_export_preserves_canonical_plan",
    "diff_tournament_projection",
    "projection_from_export_artifacts",
    "tournament_projection",
]
