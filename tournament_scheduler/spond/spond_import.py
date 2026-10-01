"""Canonical Spond bulk-import projection and workbook writer.

Owns the contract with Spond's paste/import grid. One import row represents
one *team's participation in one tournament*, so filtering a team yields
exactly one Spond event per tournament without parsing list cells. Combined
and per-team workbooks are both written from the same projection.

The Spond columns are contiguous and first, in Spond's own order. Helper
columns used for filtering/reconciliation follow them inside the same Excel
Table, so selecting the first eight columns of filtered rows pastes directly.
Cancelled tournaments are never projected as importable events.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import tempfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable, Mapping, Optional

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..models import SeasonPlan, Tournament
from ..occupancy import tournament_end_time

SPOND_IMPORT_SHEET = "Spond import"
SPOND_IMPORT_TABLE = "SpondImport"

SPOND_IMPORT_HEADERS = (
    "Startdato*",
    "Starttidspunkt",
    "Oppmøte",
    "Sluttdato",
    "Sluttidspunkt",
    "Tittel*",
    "Beskrivelse",
    "Sted",
)
HELPER_HEADERS = (
    "Lag-ID",
    "Lag",
    "Klubb",
    "Aldersgruppe",
    "Vertsklubb",
    "RVV-ID",
)

MANAGED_MANIFEST = ".spond_managed.json"

_DATE_FORMAT = "DD.MM.YYYY"
_TIME_FORMAT = "HH:MM"
_DATE_COLUMNS = (1, 4)
_TIME_COLUMNS = (2, 3, 5)


# Canonical team identity: (club, label, age group).
TeamKey = tuple[str, str, str]


@dataclass(frozen=True)
class SpondImportRow:
    """One team's participation in one non-cancelled tournament."""

    start_date: date
    start_time: Optional[time]
    meet_time: Optional[time]
    end_date: Optional[date]
    end_time: Optional[time]
    title: str
    description: str
    venue: str
    team: str
    club: str
    age_group: str
    host_club: str
    tournament_id: str

    @property
    def team_key(self) -> TeamKey:
        return (self.club, self.team, self.age_group)

    def cells(self) -> list[object]:
        return [
            self.start_date,
            self.start_time,
            self.meet_time,
            self.end_date,
            self.end_time,
            self.title,
            self.description,
            self.venue,
            team_identity(self.team_key),
            self.team,
            self.club,
            self.age_group,
            self.host_club,
            self.tournament_id,
        ]


def project_spond_import_rows(
    plan: SeasonPlan,
    ice_time_for_age_group: Optional[Mapping[str, int]] = None,
    *,
    team: TeamKey | None = None,
    club: str | None = None,
) -> list[SpondImportRow]:
    """Project the plan to deterministic per-team tournament import rows.

    *team* selects one canonical team identity ``(club, label, age_group)``.
    """
    ice_time = ice_time_for_age_group or {}
    rows: list[SpondImportRow] = []
    for tournament in sorted(plan.tournaments, key=lambda t: (t.date, t.start_time or "", t.id)):
        if tournament.cancelled:
            continue
        for participant in sorted(tournament.teams, key=lambda t: t.label):
            if team is not None and (
                participant.club,
                participant.label,
                participant.age_group,
            ) != team:
                continue
            if club is not None and participant.club != club:
                continue
            rows.append(_row_for(tournament, participant.label, participant.club, ice_time))
    return rows


def team_identity(key: TeamKey) -> str:
    """Single-cell durable team identity (club, label, age group) for filtering."""
    return " · ".join(key)


def team_keys(rows: Iterable[SpondImportRow]) -> list[TeamKey]:
    return sorted({row.team_key for row in rows})


def _row_for(
    tournament: Tournament, team_label: str, team_club: str, ice_time: Mapping[str, int]
) -> SpondImportRow:
    start = _parse_time(tournament.start_time)
    end = _parse_time(tournament_end_time(tournament, ice_time))
    return SpondImportRow(
        start_date=tournament.date,
        start_time=start,
        meet_time=None,
        end_date=tournament.date if end else None,
        end_time=end,
        title=f"{tournament.age_group} Turnering — {tournament.arena}",
        description=_description(tournament),
        venue=tournament.arena,
        team=team_label,
        club=team_club,
        age_group=tournament.age_group,
        host_club=tournament.host_club or "",
        tournament_id=tournament.id,
    )


def _description(tournament: Tournament) -> str:
    parts = [tournament.age_group]
    if tournament.host_club:
        parts.append(f"Vert: {tournament.host_club}")
    parts.append("Lag: " + ", ".join(team.label for team in tournament.teams))
    parts.append(f"RVV-ID: {tournament.id}")
    return " • ".join(parts)


def _parse_time(value: str | None) -> Optional[time]:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%H:%M").time()
    except ValueError:
        return None


def populate_spond_import_sheet(sheet, rows: list[SpondImportRow]) -> None:
    """Fill *sheet* with the import grid as a real, filterable Excel Table."""
    headers = list(SPOND_IMPORT_HEADERS + HELPER_HEADERS)
    sheet.append(headers)
    for row in rows:
        sheet.append(row.cells())

    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)
    for sheet_row in sheet.iter_rows(min_row=2):
        for column in _DATE_COLUMNS:
            sheet_row[column - 1].number_format = _DATE_FORMAT
        for column in _TIME_COLUMNS:
            sheet_row[column - 1].number_format = _TIME_FORMAT

    if rows:
        table = Table(
            displayName=SPOND_IMPORT_TABLE,
            ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        sheet.add_table(table)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    _set_widths(sheet, headers)


def _set_widths(sheet, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        longest = max(
            [len(header)]
            + [
                len(str(cell.value))
                for cell in sheet[get_column_letter(index)][1:]
                if cell.value is not None and not isinstance(cell.value, (date, time))
            ]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 60)


def write_spond_import_workbook(rows: list[SpondImportRow], output_path: str | Path) -> str:
    """Write a standalone workbook holding only the Spond import sheet."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SPOND_IMPORT_SHEET
    populate_spond_import_sheet(sheet, rows)
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))
    return str(path)


def write_per_team_workbooks(
    plan: SeasonPlan,
    output_dir: str | Path,
    ice_time_for_age_group: Optional[Mapping[str, int]] = None,
) -> dict[TeamKey, str]:
    """Reconcile ``<output_dir>`` to one workbook per participating team.

    Files written by a previous run (tracked in a manifest) that no longer
    have a projected team are removed, so a cancelled/removed team never
    leaves a stale importable workbook. Unrelated files are left alone. New
    files are built in a staging directory and moved into place afterwards.
    """
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = project_spond_import_rows(plan, ice_time_for_age_group)
    keys = team_keys(rows)
    names = _unique_filenames(keys)

    staging = Path(tempfile.mkdtemp(prefix=".staging-", dir=out_dir))
    try:
        for key in keys:
            write_spond_import_workbook(
                [row for row in rows if row.team_key == key], staging / names[key]
            )
        previous = _read_manifest(out_dir)
        for key in keys:
            os.replace(staging / names[key], out_dir / names[key])
        for stale in previous - set(names.values()):
            (out_dir / stale).unlink(missing_ok=True)
        _write_manifest(out_dir, sorted(names.values()))
    finally:
        shutil.rmtree(staging, ignore_errors=True)
    return {key: str(out_dir / names[key]) for key in keys}


def _read_manifest(out_dir: Path) -> set[str]:
    try:
        data = json.loads((out_dir / MANAGED_MANIFEST).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return set()
    return {name for name in data if isinstance(name, str) and Path(name).name == name}


def _write_manifest(out_dir: Path, names: list[str]) -> None:
    tmp = out_dir / f"{MANAGED_MANIFEST}.tmp"
    tmp.write_text(json.dumps(names, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, out_dir / MANAGED_MANIFEST)


def _unique_filenames(keys: list[TeamKey]) -> dict[TeamKey, str]:
    """Deterministic collision-safe filenames; plain label when unambiguous."""
    by_label: dict[str, list[TeamKey]] = {}
    for key in keys:
        by_label.setdefault(_safe_filename(key[1]).casefold(), []).append(key)
    names: dict[TeamKey, str] = {}
    used: set[str] = set()
    for group in by_label.values():
        for key in group:
            club, label, age_group = key
            base = _safe_filename(label)
            if len(group) > 1:
                base = f"{base} ({_safe_filename(age_group)}, {_safe_filename(club)})"
            name, n = f"{base}.xlsx", 2
            while name.casefold() in used:
                name, n = f"{base} {n}.xlsx", n + 1
            used.add(name.casefold())
            names[key] = name
    return names


def _safe_filename(label: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', "_", label).strip(" .")
    return cleaned or "lag"
